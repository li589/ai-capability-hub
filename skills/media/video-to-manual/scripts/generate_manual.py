# -*- coding: utf-8 -*-
"""
Reusable Word document generator for operation manuals.
Accepts a JSON config file that defines the document structure,
embeds screenshots with proper formatting.

Usage:
    python generate_manual.py <config.json> [-o output.docx] [--screenshot-dir ./screenshots]
"""

import json, os, sys, argparse
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

FONT_NAME = '微软雅黑'
FONT_SIZE = 11
SCREENSHOT_WIDTH = Inches(5.5)


class ManualGenerator:
    def __init__(self, screenshot_dir='.'):
        self.doc = Document()
        self.screenshot_dir = screenshot_dir
        self._setup_styles()

    def _setup_styles(self):
        style = self.doc.styles['Normal']
        font = style.font
        font.name = FONT_NAME
        font.size = Pt(FONT_SIZE)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)

    def _set_font(self, run, name=FONT_NAME, size=FONT_SIZE):
        run.font.name = name
        run.font.size = Pt(size)
        run.element.rPr.rFonts.set(qn('w:eastAsia'), name)

    def add_title_page(self, title, subtitle='', date_line=''):
        """Add centered title page."""
        t = self.doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = t.add_run(title)
        r.bold = True
        self._set_font(r, size=24)
        r.font.color.rgb = RGBColor(0, 51, 102)

        if subtitle:
            s = self.doc.add_paragraph()
            s.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = s.add_run(subtitle)
            self._set_font(r, size=13)
            r.font.color.rgb = RGBColor(89, 89, 89)

        if date_line:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(date_line)
            self._set_font(r, size=10)
            r.font.color.rgb = RGBColor(128, 128, 128)

        self.doc.add_paragraph()
        self.doc.add_paragraph()

    def add_heading(self, text, level=1):
        h = self.doc.add_heading(text, level=level)
        for run in h.runs:
            run.font.name = FONT_NAME
            run.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)
        return h

    def add_para(self, text, bold=False, size=FONT_SIZE, color=None, align=None):
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        r.bold = bold
        self._set_font(r, size=size)
        if color:
            r.font.color.rgb = RGBColor(*color)
        if align:
            p.alignment = align
        return p

    def add_bullet(self, text, level=0):
        p = self.doc.add_paragraph(text, style='List Bullet')
        p.paragraph_format.left_indent = Cm(1.5 + level * 0.8)
        for run in p.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(FONT_SIZE)
            run.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)
        return p

    def add_numbered(self, text):
        p = self.doc.add_paragraph(text, style='List Number')
        for run in p.runs:
            run.font.name = FONT_NAME
            run.font.size = Pt(FONT_SIZE)
            run.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)
        return p

    def add_code(self, text):
        p = self.doc.add_paragraph()
        r = p.add_run(text)
        self._set_font(r, name='Consolas', size=10)
        p.paragraph_format.left_indent = Cm(1)
        return p

    def add_note(self, text, color=(0, 102, 204)):
        p = self.doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = FONT_NAME
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(*color)
        run.bold = True
        run.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)
        p.paragraph_format.left_indent = Cm(1)
        return p

    def add_screenshot(self, filename, caption='', width=SCREENSHOT_WIDTH):
        filepath = os.path.join(self.screenshot_dir, filename)
        if not os.path.exists(filepath):
            self.add_para(f'[截图缺失: {filename}]', color=(255, 0, 0))
            return None

        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        try:
            r.add_picture(filepath, width=width)
        except Exception as e:
            self.add_para(f'[图片加载失败: {filename} - {e}]', color=(255, 0, 0))
            return None

        if caption:
            cap = self.doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(caption)
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(128, 128, 128)
            r.font.name = FONT_NAME
            r.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)

        self.doc.add_paragraph()
        return filepath

    def add_table(self, headers, rows):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = 'Light Grid Accent 1'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            cell = table.rows[0].cells[i]
            cell.text = h
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
                    run.font.name = FONT_NAME
                    run.font.size = Pt(10)
                    run.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                cell = table.rows[r_idx + 1].cells[c_idx]
                cell.text = str(val)
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = FONT_NAME
                        run.font.size = Pt(10)
                        run.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_NAME)
        self.doc.add_paragraph()
        return table

    def add_page_break(self):
        self.doc.add_page_break()

    def add_element(self, el):
        """Add an element from the JSON config."""
        t = el.get('type', '')
        if t == 'heading':
            self.add_heading(el['text'], el.get('level', 1))
        elif t == 'para':
            self.add_para(
                el['text'],
                bold=el.get('bold', False),
                size=el.get('size', FONT_SIZE),
                color=tuple(el['color']) if el.get('color') else None,
            )
        elif t == 'bullet':
            if 'items' in el:
                for item in el['items']:
                    self.add_bullet(item, el.get('level', 0))
            else:
                self.add_bullet(el['text'], el.get('level', 0))
        elif t == 'numbered':
            if 'items' in el:
                for item in el['items']:
                    self.add_numbered(item)
            else:
                self.add_numbered(el['text'])
        elif t == 'code':
            self.add_code(el['text'])
        elif t == 'note':
            if el.get('color'):
                self.add_note(el['text'], color=tuple(el['color']))
            else:
                self.add_note(el['text'])
        elif t == 'screenshot':
            self.add_screenshot(el['file'], el.get('caption', ''))
        elif t == 'table':
            self.add_table(el['headers'], el['rows'])
        elif t == 'page_break':
            self.add_page_break()
        elif t == 'blank':
            self.doc.add_paragraph()

    def save(self, path):
        self.doc.save(path)
        size_kb = os.path.getsize(path) / 1024
        print(f"Document saved: {path} ({size_kb:.0f} KB)")
        return path


def main():
    parser = argparse.ArgumentParser(description='Generate Word manual from JSON config')
    parser.add_argument('config', help='JSON config file path')
    parser.add_argument('-o', '--output', default=None, help='Output .docx path')
    parser.add_argument('--screenshot-dir', default='.', help='Screenshot directory')
    args = parser.parse_args()

    with open(args.config, 'r', encoding='utf-8') as f:
        cfg = json.load(f)

    output_path = args.output or cfg.get('output', 'output.docx')

    gen = ManualGenerator(screenshot_dir=args.screenshot_dir)

    # Title page
    gen.add_title_page(
        cfg.get('title', '操作手册'),
        cfg.get('subtitle', ''),
        cfg.get('date_line', ''),
    )

    # Generate sections
    sections = cfg.get('sections', [])
    for section in sections:
        if isinstance(section, dict):
            gen.add_element(section)

    gen.save(output_path)


if __name__ == '__main__':
    main()
