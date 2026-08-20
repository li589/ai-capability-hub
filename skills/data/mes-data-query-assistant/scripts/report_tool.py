# -*- coding: utf-8 -*-
"""
报告生成工具：将一次问数的会话结果（多轮查询的 question + parse + results + 图表）汇总为专业报告。

默认输出 Word（.docx），可选 --html 输出 HTML（浏览器直接展示，内嵌图表 base64）。

输入：一个 JSON 文件（可由 query_tool 多轮输出拼接，或人工构造），结构：
[
  {"question": "今日产量", "parse": {...}, "results": [{"tag":"...", "sql":"...", "data":[{"列":"值"}], "warns":[]}], "summary": "人工/AI 结论", "chart": "reports/charts/xxx.png"},
  ...
]

用法（必须用 venv python 跑，含 python-docx）：
  python report_tool.py sessions.json -o 报告.docx                # Word 报告
  python report_tool.py sessions.json -o 报告.html --html         # HTML 报告
  python report_tool.py sessions.json --html                      # 默认 reports/智能问数报告.html

报告结构：
  封面（标题/日期/数据源）→ 摘要（各问题结论）→ 明细（每轮：问题/解析/查询/表格/图表/注意点）→ 附录（SQL）
"""
import sys, os, json, base64, datetime, mimetypes

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
REPORTS_DIR = os.path.join(BASE, "reports")
DEFAULT_OUT = os.path.join(REPORTS_DIR, "智能问数报告.docx")


def _load_sessions(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def _img_b64(path):
    if not path or not os.path.exists(path):
        return None
    mime = mimetypes.guess_type(path)[0] or "image/png"
    with open(path, "rb") as f:
        return f"data:{mime};base64," + base64.b64encode(f.read()).decode()


# ---------------- Word ----------------
def build_word(sessions, out_path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    # 封面标题
    t = doc.add_heading("MES 智能问数分析报告", 0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"生成时间: {datetime.date.today()}").font.size = Pt(11)
    p.add_run(f"\n数据源: MySQL 只读连接").font.size = Pt(11)

    # 摘要
    doc.add_heading("一、摘要", level=1)
    for s in sessions:
        q = s.get("question", "")
        summ = s.get("summary") or s.get("conclusion") or ""
        key = s.get("key_finding") or s.get("problem") or ""
        doc.add_paragraph(f"• {q}")
        if summ:
            doc.add_paragraph(f"    结论: {summ}", style="List Bullet")
        if key:
            doc.add_paragraph(f"    注意: {key}", style="List Bullet")

    # 明细
    doc.add_heading("二、明细", level=1)
    for i, s in enumerate(sessions, 1):
        doc.add_heading(f"{i}. {s.get('question', '未命名问题')}", level=2)
        pr = s.get("parse") or {}
        if pr:
            doc.add_paragraph(f"解析: " + " | ".join(f"{k}={v}" for k, v in pr.items()))
        for r in s.get("results", []):
            if r.get("tag"):
                doc.add_paragraph(f"[{r['tag']}]", style="List Bullet")
            if r.get("data"):
                _doc_table(doc, r["data"])
            if r.get("warns"):
                for w in r["warns"]:
                    doc.add_paragraph(f"⚠️ {w}", style="List Bullet")
        if s.get("summary"):
            doc.add_paragraph(f"结论: {s['summary']}")
        if s.get("key_finding"):
            doc.add_paragraph(f"⚠️ 注意点: {s['key_finding']}")
        if s.get("possible_causes"):
            doc.add_paragraph("可能原因:")
            for c in s["possible_causes"]:
                doc.add_paragraph(f"  • {c}", style="List Bullet")
        if s.get("followup_questions"):
            doc.add_paragraph("可追问:")
            for fq in s["followup_questions"]:
                doc.add_paragraph(f"  • {fq}", style="List Bullet")
        chart = s.get("chart")
        if chart and os.path.exists(chart):
            doc.add_picture(chart, width=Inches(5.5))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 附录 SQL
    doc.add_heading("三、附录：查询 SQL", level=1)
    for s in sessions:
        for r in s.get("results", []):
            if r.get("sql"):
                doc.add_paragraph(f"{s.get('question')} / {r.get('tag')}:")
                doc.add_paragraph(r["sql"])

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    return out_path


def _doc_table(doc, data):
    if not data:
        return
    cols = list(data[0].keys())
    rows = [list(d.values()) for d in data]
    table = doc.add_table(rows=1 + len(rows), cols=len(cols))
    table.style = "Light Grid Accent 1"
    for j, c in enumerate(cols):
        table.rows[0].cells[j].text = str(c)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            table.rows[i].cells[j].text = str(v)


# ---------------- HTML ----------------
def build_html(sessions, out_path):
    cards = []
    for i, s in enumerate(sessions, 1):
        pr = s.get("parse") or {}
        parse_html = " | ".join(f"<b>{k}</b>={v}" for k, v in pr.items()) if pr else ""
        results_html = []
        for r in s.get("results", []):
            h = [f"<h4>📊 {r.get('tag','')}</h4>"] if r.get("tag") else []
            if r.get("data"):
                cols = list(r["data"][0].keys()) if r["data"] else []
                h.append("<table><thead><tr>" + "".join(f"<th>{c}</th>" for c in cols) + "</tr></thead><tbody>")
                for d in r["data"]:
                    h.append("<tr>" + "".join(f"<td>{v}</td>" for v in d.values()) + "</tr>")
                h.append("</tbody></table>")
            for w in r.get("warns", []):
                h.append(f'<div class="warn">⚠️ {w}</div>')
            results_html.append("".join(h))
        chart = s.get("chart")
        img_html = f'<img src="{_img_b64(chart)}" class="chart">' if chart and _img_b64(chart) else ""
        causes = "".join(f"<li>{c}</li>" for c in s.get("possible_causes", []))
        follow = "".join(f"<li>{fq}</li>" for fq in s.get("followup_questions", []))
        cards.append(f"""
        <div class="card">
          <h3>{i}. {s.get('question','')}</h3>
          <div class="meta">{parse_html}</div>
          {img_html}
          {"".join(results_html)}
          <div class="summary">{s.get('summary','')}</div>
          {"<div class='finding'>🔍 注意点: "+s['key_finding']+"</div>" if s.get("key_finding") else ""}
          {"<div class='causes'><b>可能原因:</b><ul>"+causes+"</ul></div>" if causes else ""}
          {"<div class='follow'><b>可追问:</b><ul>"+follow+"</ul></div>" if follow else ""}
        </div>""")

    html = f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>MES 智能问数分析报告</title>
<style>
  body {{ font-family: "Microsoft YaHei", sans-serif; margin: 0; background: #f4f6f9; color: #222; }}
  .wrap {{ max-width: 960px; margin: 0 auto; padding: 24px; }}
  .hero {{ background: linear-gradient(135deg, #1a3a6b, #2d5aa8); color: #fff; border-radius: 12px; padding: 28px; margin-bottom: 20px; }}
  .hero h1 {{ margin: 0 0 6px; }}
  .hero p {{ margin: 2px 0; opacity: .85; }}
  .card {{ background: #fff; border-radius: 10px; padding: 20px; margin-bottom: 18px; box-shadow: 0 1px 4px rgba(0,0,0,.08); }}
  .card h3 {{ margin: 0 0 8px; color: #1a3a6b; }}
  .meta {{ color: #777; font-size: 13px; margin-bottom: 10px; }}
  table {{ border-collapse: collapse; width: 100%; margin: 10px 0; font-size: 13px; }}
  th, td {{ border: 1px solid #dfe3ea; padding: 6px 10px; text-align: left; }}
  th {{ background: #eef2f8; }}
  tr:nth-child(even) td {{ background: #fafbfd; }}
  .chart {{ max-width: 100%; border-radius: 8px; margin: 8px 0; }}
  .warn {{ background: #fff7e6; border-left: 4px solid #faad14; padding: 8px 12px; margin: 8px 0; border-radius: 4px; font-size: 13px; }}
  .summary {{ background: #f0f7ff; border-left: 4px solid #1890ff; padding: 10px 14px; margin: 10px 0; border-radius: 4px; }}
  .finding {{ background: #fff1f0; border-left: 4px solid #ff4d4f; padding: 8px 12px; margin: 8px 0; border-radius: 4px; }}
  .causes, .follow {{ margin: 8px 0; font-size: 14px; }}
</style></head><body><div class="wrap">
  <div class="hero"><h1>MES 智能问数分析报告</h1>
    <p>生成时间: {datetime.date.today()} ｜ 数据源: MySQL 只读连接 ｜ 问题数: {len(sessions)}</p></div>
  {''.join(cards)}
</div></body></html>"""
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)
    return out_path


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        return
    sessions_path = args[0]
    out_path = DEFAULT_OUT
    as_html = False
    for i, a in enumerate(args):
        if a == "-o" and i + 1 < len(args):
            out_path = args[i + 1]
        elif a == "--html":
            as_html = True
            if out_path == DEFAULT_OUT:
                out_path = os.path.join(REPORTS_DIR, "智能问数报告.html")
    try:
        sessions = _load_sessions(sessions_path)
    except Exception as e:
        print(f"❌ 读取会话文件失败: {e}")
        return
    try:
        if as_html:
            path = build_html(sessions, out_path)
        else:
            path = build_word(sessions, out_path)
        print(path)
    except ImportError:
        print("❌ 缺少 python-docx，请先安装: pip install python-docx")
    except Exception as e:
        print(f"❌ 报告生成失败: {e}")


if __name__ == "__main__":
    main()
