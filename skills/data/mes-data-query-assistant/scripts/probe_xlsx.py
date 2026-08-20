#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""探查 schema 3.0.xlsx 结构：sheet 列表 + 每个 sheet 前 5 行"""
import openpyxl, sys

SRC = r"{{HOME}}/Desktop/schema 3.0 .xlsx"
wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n===== sheet: {sn}  dims={ws.max_row}x{ws.max_column} =====")
    cnt = 0
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        vals = [str(v)[:40] if v is not None else "" for v in row]
        print(" | ".join(vals))
        cnt += 1
        if cnt >= 5:
            break
wb.close()
