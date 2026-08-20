#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
企微客户触达编排 · 四件套 Excel 空白模板生成器
照搬技能四件套 Markdown 结构成 4 个 Sheet：触点编排矩阵 / 话术填充表 / 频次排程表 / 合规折叠校验清单。

纯本地、免 Key。依赖 openpyxl；缺依赖时退出码 2 并给安装引导，--install-deps 可自动装。
三级兜底写盘：--out 指定路径 → 当前工作目录 → /tmp。成功退出码 0，
stdout 协议行：REACH_XLSX_FILE=<绝对路径> 供 agent 解析。
"""

import sys
import os
import subprocess

XLSX_FILENAME = "企微客户触达编排四件套模板.xlsx"

# ---------- 依赖检测 ----------
def ensure_openpyxl(auto_install=False):
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        sys.stderr.write(
            "\n[缺依赖] 本脚本需要 openpyxl 来生成 Excel。\n"
            "安装方法（任选其一）：\n"
            "  1) python3 -m pip install openpyxl\n"
            "  2) python3 {script} --install-deps   # 本脚本自动安装\n"
            "装好后重新运行本脚本即可。\n\n".format(script=sys.argv[0])
        )
        if auto_install:
            sys.stderr.write("[自动安装] 正在执行: python3 -m pip install openpyxl ...\n")
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
                return True
            except Exception as e:  # noqa: BLE001
                sys.stderr.write("[自动安装失败] {}\n请手动安装：python3 -m pip install openpyxl\n".format(e))
        sys.exit(2)


# ---------- 写盘路径解析（三级兜底）----------
def resolve_out_path(out_arg):
    if out_arg:
        out_dir = os.path.dirname(os.path.abspath(out_arg)) or "."
        if os.access(out_dir if os.path.isdir(out_dir) else os.path.dirname(out_dir) or os.getcwd(), os.W_OK):
            return os.path.abspath(out_arg)
        # 指定路径不可写，降级
    # 兜底 1：当前工作目录
    try:
        cwd = os.getcwd()
        test = os.path.join(cwd, ".wb_reach_writetest")
        with open(test, "w") as f:
            f.write("x")
        os.remove(test)
        return os.path.join(cwd, XLSX_FILENAME)
    except OSError:
        pass
    # 兜底 2：/tmp
    return os.path.join("/tmp", XLSX_FILENAME)


# ---------- Sheet 定义 ----------
SHEET_TOUCHPOINT_HEADERS = [
    "人群(价值档×阶段)", "主触点", "备触点", "触达目标",
    "适配判据(精准度×覆盖×成本×额度)", "频次上限(按价值量化)",
]
SHEET_TOUCHPOINT_NOTES = [
    "触点选项：1v1私聊 / 朋友圈 / 群发 / 社群（可组合，但须频次总控防重复打扰）",
    "适配判据四维：1v1精准高成本高；朋友圈广覆盖低成本但额度最稀缺(3/天4/月)；群发广覆盖低成本但额度共用(企业+个人共享,单次200人)；社群中可控但易退群",
    "选触点：高价值且目标转化/召回→1v1；长尾促活→朋友圈/群发；高频低客单→社群1v多(1v1不划算)",
    "价值档(引用RFM不重算)：高/中/低；阶段(引用生命周期旅程)：引入期/成长期/成熟期/休眠期/流失期",
    "频次必须按价值量化(如1v1月≤2次主动推送)，禁「适度」模糊词",
]

SHEET_SCRIPT_HEADERS = [
    "人群", "触点", "钩子(利益/痛点)", "价值点(一条)",
    "CTA(明确动作)", "素材(图·卡片)", "可替代表达", "不可使用边界",
]
SHEET_SCRIPT_NOTES = [
    "单条≤100字一核心点；钩子+价值点+CTA+素材四要素齐全",
    "B2B重案例/方法论背书；B2C重人设生活+信任——按行业倾向填，不通用一刀切",
    "群发文案≤4000字+附件≤9；禁广告法敏感词(最/第一/国家级)，换「较受欢迎」类合规表达",
    "话术按人群状态×触点×任务组织，非按欢迎语/活动文案平铺",
]

SHEET_FREQ_HEADERS = [
    "日期·时段", "人群", "触点", "内容主题",
    "额度预算(本月剩多少)", "额度够否", "备注",
]
SHEET_FREQ_NOTES = [
    "额度假约束(★已核验·按最新规则核验需复核)：",
    "  朋友圈：每客每自然日≤3条成员发表 + 每客每月≤4条企业发表",
    "  群发(旧规)：每客每日1条+每月4条企业(单次200人)；新规可配三档(每日1/每周7/当月天数条)",
    "  个人群发与企业群发共享同一额度池——企业用了企业群发会消耗个人群发额度",
    "  超上限：客户收不到(发送方不报错)",
    "同人群同窗口期无跨触点重复打扰；有时按价值密度合并去重留主触点",
    "时段选用户决策窗口(通勤/午休/晚浏览高峰·行业经验值非实证)，非纯定时盲发",
]

SHEET_COMPLIANCE_HEADERS = [
    "触达动作", "风险项", "风险等级(高/中/低)", "修正建议",
]
SHEET_COMPLIANCE_NOTES = [
    "校验项：①敏感词(广告法最/第一/国家级) ②折叠风险(文案>4000字/附件>9/纯链接无文案)",
    "  ③超额度(朋友圈>3/天4/月、群发超共享额度) ④2026.7新规(外挂零容忍/频率管控/拉群需确认)",
    "  ⑤反模式(宣称无限群发/无需确认自动群发/外挂绕过)",
    "触及反模式→标⚠要求改手动1v1或合规群发；平台规则标【按最新规则核验·需复核】提示用户复核官方",
]


# ---------- 构建 Excel ----------
def build_workbook():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F5496")
    note_font = Font(italic=True, color="666666")
    wrap = Alignment(wrap_text=True, vertical="top")

    sheets = [
        ("①触点编排矩阵", SHEET_TOUCHPOINT_HEADERS, SHEET_TOUCHPOINT_NOTES),
        ("②话术填充表", SHEET_SCRIPT_HEADERS, SHEET_SCRIPT_NOTES),
        ("③频次排程表", SHEET_FREQ_HEADERS, SHEET_FREQ_NOTES),
        ("④合规折叠校验清单", SHEET_COMPLIANCE_HEADERS, SHEET_COMPLIANCE_NOTES),
    ]

    # 默认第一个 sheet 复用，其余新建
    for idx, (sheet_name, headers, notes) in enumerate(sheets):
        ws = wb.worksheets[0] if idx == 0 else wb.create_sheet()
        ws.title = sheet_name

        # 备注行（写在表头上方）
        for i, note in enumerate(notes, start=1):
            cell = ws.cell(row=i, column=1, value=note)
            cell.font = note_font
            cell.alignment = wrap
        header_row = len(notes) + 2  # 留一空行

        # 表头
        for col, h in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap

        # 预留若干空行
        for r in range(header_row + 1, header_row + 16):
            for col in range(1, len(headers) + 1):
                ws.cell(row=r, column=col).alignment = wrap

        # 列宽
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22

        ws.sheet_view.showGridLines = True

    return wb


# ---------- 入口 ----------
def main():
    auto_install = "--install-deps" in sys.argv
    out_arg = None
    for a in sys.argv[1:]:
        if a.startswith("--out="):
            out_arg = a[len("--out="):]
        elif a not in ("--install-deps", "-h", "--help"):
            if not a.startswith("-") and out_arg is None:
                out_arg = a

    if "-h" in sys.argv or "--help" in sys.argv:
        print(__doc__)
        print("\n用法：python3 gen_reach_xlsx.py [--out=路径.xlsx] [--install-deps]")
        sys.exit(0)

    ensure_openpyxl(auto_install=auto_install)
    wb = build_workbook()
    out_path = resolve_out_path(out_arg)
    wb.save(out_path)
    print("REACH_XLSX_FILE=" + os.path.abspath(out_path))
    sys.exit(0)


if __name__ == "__main__":
    main()
