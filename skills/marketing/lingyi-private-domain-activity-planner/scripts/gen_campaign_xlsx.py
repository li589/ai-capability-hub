#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
零一数科 · 内容与互动活动 — 四件套 Excel 模板生成器
生成 5-Sheet 空白模板（表头/冻结首行/列宽/示例行就位），拿来即填。

用法:
  python3 gen_campaign_xlsx.py --out ./内容与互动活动模板.xlsx
  python3 gen_campaign_xlsx.py                 # 默认写到当前目录
  python3 gen_campaign_xlsx.py --install-deps  # 缺 openpyxl 时自动安装后重试

工程规范:
  - 纯本地、不联网（仅 pip 安装 openpyxl 时联网）、不扣点
  - 三级兜底写盘: --out → 当前目录 → /tmp
  - 退出码: 0=成功, 2=缺依赖(给安装引导)
  - 成功 stdout 协议行: CAMPAIGN_XLSX_FILE=<绝对路径>  供 agent 解析
"""

import sys
import os

SHEETS = [
    "使用说明",
    "月度内容排期",
    "活动方案",
    "玩法机制",
    "物料BOM",
]


def need_openpyxl():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        return False


def install_deps():
    print("检测到缺少 openpyxl，正在自动安装...", file=sys.stderr)
    import subprocess
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
        return True
    except Exception as e:  # noqa: BLE001
        print(f"自动安装失败: {e}", file=sys.stderr)
        return False


def build_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="2F5496")
    note_font = Font(name="微软雅黑", size=10, italic=True, color="595959")
    cell_font = Font(name="微软雅黑", size=10)
    example_font = Font(name="微软雅黑", size=10, color="808080", italic=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def sheet_with_header(title, headers, widths, examples=None, notes=None):
        ws = wb.create_sheet(title)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for ci, h in enumerate(headers, start=1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center
            c.border = border
        ws.row_dimensions[1].height = 28
        ws.freeze_panes = "A2"
        if examples:
            for ri, row in enumerate(examples, start=2):
                for ci, val in enumerate(row, start=1):
                    c = ws.cell(row=ri, column=ci, value=val)
                    c.font = example_font
                    c.alignment = left
                    c.border = border
                ws.row_dimensions[ri].height = 22
        if notes:
            base = (len(examples) + 2) if examples else 2
            for ri, note in enumerate(notes, start=base):
                ws.cell(row=ri, column=1, value=note).font = note_font
                ws.merge_cells(start_row=ri, start_column=1, end_row=ri, end_column=len(headers))
                ws.row_dimensions[ri].height = 20
        return ws

    # === Sheet 1: 使用说明 ===
    ws1 = wb.create_sheet("使用说明")
    ws1.column_dimensions["A"].width = 22
    ws1.column_dimensions["B"].width = 90
    ws1.freeze_panes = "A2"
    guide = [
        ("项", "说明"),
        ("四件套", "①月度内容排期 ②活动方案 ③玩法机制 ④物料BOM，逐项填空，对话产出与之对应"),
        ("4321内容配比", "干货40% / 生活人设30% / 客户证言20% / 促销种草10%；广告>30%易屏蔽"),
        ("内容配比判定", "干货≥35%且促销≤15%通过；促销>30%→标反模式要求调比例"),
        ("栏目完整判定", "固定栏目≥3类(干货/证言/互动)且周覆盖通过；只促销无日常→补日常栏目"),
        ("发布节奏", "互动型白天推送、促销型晚上推送；早8干货/午12互动/晚8种草，间隔≥2h"),
        ("K系数自检(裂变)", "K=人均邀请数×邀请转化率；K≥1可放量，K<1调激励/门槛/钩子；禁高/中/低模糊词"),
        ("非裂变玩法", "打卡/答题/UGC晒单 K不适用，改用任务完成率/UGC产出率作主指标"),
        ("活跃→转化咬合", "每个活动方案必含活跃→转化映射(如互动率→复购券触发)；缺=为活跃而活跃反模式"),
        ("物料BOM判定", "预算合计≤总预算通过；关键物料到货≤活动上线；海报四要素齐(主标/副标/行动指令/紧迫感)"),
        ("权限分级", "对外版隐藏预算字段，内部版全字段；各供应商互相不可见"),
        ("合规-UGC", "审核(敏感词/违禁图)+肖像授权(先授权后使用)+未成年人保护(不满10周岁不得代言)"),
        ("合规-抽奖5条", "明示条件/不虚构奖品/奖金不超限/每局每日押输赢上限/联网游戏版号备案"),
        ("合规-营销", "禁诱导分享(转发得奖励→改参与解锁)、禁虚构原价、禁广告法极限词(最/第一/国家级)"),
        ("防刷四闸", "同设备IP限+奖励延迟兑现(24h未退群)+风控阈值+仅新用户计邀请"),
        ("待印证/校准标", "行业基准用区间标【待印证·按业务校准】；平台规则标【按2026规则核验·需复核】；节日标【以当年官方公告为准】"),
    ]
    for ri, (k, v) in enumerate(guide, start=1):
        a = ws1.cell(row=ri, column=1, value=k)
        b = ws1.cell(row=ri, column=2, value=v)
        if ri == 1:
            a.font = header_font; a.fill = header_fill; a.alignment = center
            b.font = header_font; b.fill = header_fill; b.alignment = center
        else:
            a.font = Font(name="微软雅黑", bold=True, size=10); a.alignment = left
            b.font = cell_font; b.alignment = left
        a.border = border; b.border = border
        ws1.row_dimensions[ri].height = 24 if ri == 1 else 30

    # === Sheet 2: 月度内容排期（12列）===
    sheet_with_header(
        "月度内容排期",
        headers=["周次", "日期", "时段(早8/午12/晚8)", "栏目类型(干货/生活人设/客证言/促销种草)",
                 "内容主题", "适用人群分层", "触达渠道(朋友圈/社群/公众号/私聊)", "文案钩子",
                 "负责人", "素材就绪(待制作/制作中/就绪)", "发布(待发/已发)", "复盘指标(打开/互动/转化)"],
        widths=[8, 14, 16, 24, 20, 16, 22, 22, 10, 16, 12, 22],
        examples=[["W1", "2026-09-02", "早8", "干货", "秋冬5款茶底科普", "全员", "社群", "你喝的是哪种茶底？评论区测",
                   "小林", "就绪", "待发", "打开X%·互动Y·转化Z"]],
        notes=["配比自检: 干货[ ]%≥35% | 促销[ ]%≤15% | 广告合计[ ]%<30% | 固定栏目[ ]类≥3且周覆盖"],
    )

    # === Sheet 3: 活动方案 ===
    sheet_with_header(
        "活动方案",
        headers=["章节", "字段", "内容"],
        widths=[16, 26, 70],
        examples=[
            ["6要素", "名称", "[活动名]"],
            ["6要素", "主题&口号", "[主题/口号，与目标一致·引起共鸣·简洁易记]"],
            ["6要素", "周期", "[预热D-3~D-1 / 正式D-day / 收尾3h]"],
            ["6要素", "目标(含活跃→转化咬合)", "[拉新/促活/转化；咬合映射如互动率→复购率]"],
            ["6要素", "人群", "[新/老用户、分层、规模]"],
            ["6要素", "总预算", "[金额]"],
            ["预热D-3~D-1", "目标/动作/话术/指标", "[建信任；痛点种草→价值输出→悬念剧透；7日互动>40%蓄水]"],
            ["正式D-day", "目标/动作/话术/指标", "[引爆成交；开场→方案→晒单从众→锁单；转化率/客单/接龙]"],
            ["收尾3h", "目标/动作/话术/指标", "[逼单补单；倒计时3h/1h/30min→售罄→补单；收尾转化占比]"],
            ["奖励机制", "类型/数量/上限/发放", "[物质×社交货币×身份；每用户上限/总量/连续；自动/手动+时效]"],
            ["传播出口", "一级/二级/裂变钩子", "[海报+群内卡+邀请链接 / 小程序卡+直播+公众号 / 邀X人解锁]"],
            ["人员分工", "角色/任务", "[内容/活动/社群/设计 分工至个人]"],
            ["风险合规", "合规/防刷/熔断", "[抽奖5条/UGC审核授权未成年/无诱导分享|四闸|K>1.5熔断]"],
            ["复盘指标", "GMV/客单/ROI/咬合验证", "[总转化率/各段流失/活跃→转化咬合验证]"],
        ],
        notes=["必含活跃→转化咬合映射，缺=为活跃而活跃反模式(异常6)；合规缺项触发异常1-3"],
    )

    # === Sheet 4: 玩法机制 ===
    sheet_with_header(
        "玩法机制",
        headers=["要素", "内容"],
        widths=[24, 80],
        examples=[
            ["玩法类型", "[任务宝/助力/拼团/分销/UGC征集/打卡/答题/抽奖/榜单]"],
            ["选型依据", "[低客单助力拼团/中客单任务宝/高客单分销/日常打卡答题UGC]"],
            ["参与门槛", "[如发图+话题/满299/邀1人]"],
            ["任务动作", "[触达→参与→行动→兑奖→(分享)]"],
            ["奖励兑现", "[领券/抽奖/积分到账+上限/时效]"],
            ["传播出口", "[朋友圈/群内分享/邀请好友/海报/小程序卡]"],
            ["门槛三阶梯", "基础档[拉参与]/进阶档[拉裂变]/头部档[拉传播]"],
            ["激励组合", "[物质(实物/券)×社交货币(排行榜)×身份(进阶群/勋章)]"],
            ["K系数自检(裂变)", "K=人均邀请[X]×转化率[Y]=[K值]；≥1放量/<1调激励门槛(非裂变标K不适用改任务完成率)"],
            ["防刷四闸", "1同设备IP限 2延迟兑现24h 3风控阈值 4仅新用户计邀请"],
        ],
        notes=["裂变玩法K必给数值≥1；非裂变用任务完成率/UGC产出率；防刷四闸缺触发异常4"],
    )

    # === Sheet 5: 物料BOM（10列）===
    sheet_with_header(
        "物料BOM",
        headers=["物料名称", "规格", "数量", "用途(对应方案节点)", "渠道", "负责人",
                 "到货·上线时间", "预算", "状态(待采/制作中/就绪)", "权限(对内/对外)"],
        widths=[18, 16, 10, 22, 14, 10, 16, 12, 16, 14],
        examples=[["会员日主海报", "750×1334", "4张/月", "预热+正式", "朋友圈+社群", "设计",
                   "2026-08-30", "1000", "就绪", "对外"]],
        notes=["预算合计[ ]≤总预算[ ](超触发异常5) | 关键物料到货≤活动上线(晚触发异常5) | 海报四要素:主标+副标+行动指令+紧迫感 | 对外版隐藏预算字段"],
    )

    # 删默认 sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    return wb


def write_out(wb, out_path):
    # 三级兜底: --out → 当前目录 → /tmp
    candidates = []
    if out_path:
        candidates.append(os.path.abspath(out_path))
    candidates.append(os.path.abspath("内容与互动活动模板.xlsx"))
    candidates.append(os.path.join("/tmp", "内容与互动活动模板.xlsx"))

    last_err = None
    for path in candidates:
        try:
            dirname = os.path.dirname(path)
            if dirname and not os.path.exists(dirname):
                os.makedirs(dirname, exist_ok=True)
            wb.save(path)
            return path
        except Exception as e:  # noqa: BLE001
            last_err = e
            continue
    raise RuntimeError(f"三级兜底写盘均失败: {last_err}")


def main():
    import argparse
    parser = argparse.ArgumentParser(description="内容与互动活动四件套 Excel 模板生成器")
    parser.add_argument("--out", help="输出 xlsx 路径；缺省走三级兜底(--out→当前目录→/tmp)")
    parser.add_argument("--install-deps", action="store_true", help="缺 openpyxl 时自动安装后重试")
    args = parser.parse_args()

    if not need_openpyxl():
        if args.install_deps:
            if not install_deps():
                print("缺少依赖 openpyxl，自动安装失败。请手动安装: pip3 install openpyxl", file=sys.stderr)
                sys.exit(2)
            if not need_openpyxl():
                print("openpyxl 安装后仍不可用，请检查 Python 环境。", file=sys.stderr)
                sys.exit(2)
        else:
            print("缺少依赖 openpyxl。请安装后重试：", file=sys.stderr)
            print("  pip3 install openpyxl", file=sys.stderr)
            print("或带 --install-deps 自动安装后再跑：", file=sys.stderr)
            print(f"  {sys.argv[0]} --install-deps", file=sys.stderr)
            sys.exit(2)

    try:
        wb = build_workbook()
        path = write_out(wb, args.out)
        print(f"CAMPAIGN_XLSX_FILE={path}")
        print(f"Excel 模板已生成：{path}，打开即可填；对话里的四件套 Markdown 与 Sheets 一一对应。", file=sys.stderr)
        sys.exit(0)
    except Exception as e:  # noqa: BLE001
        print(f"生成失败: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
