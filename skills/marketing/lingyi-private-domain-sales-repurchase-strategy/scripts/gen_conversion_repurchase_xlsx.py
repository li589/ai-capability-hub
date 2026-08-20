#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""转化与复购运营 —— Excel 模板生成脚本 (v0.1.0)

生成一份「拿来即填」的转化与复购运营空白 Excel 模板（含表头样式 / 冻结 / 列宽 / 示例行），
便于运营岗落盘后直接在 Excel 里填自己的首购转化漏斗、复购策略表、企微话术组、沉睡召回方案。

五个 Sheet：
  1. 使用说明
  2. 首购转化漏斗（场景 × 五要素）
  3. 复购策略表（复购品类 × 补货周期/触发/推荐/权益/话术）
  4. 企微话术组（话术ID × 五段式/触发/人设/答疑FAQ）
  5. 沉睡召回方案（召回次序 × 分层/触达/权益/归因）

依赖：openpyxl（缺失时本脚本引导安装，不强制）。
纯本地、不联网、不扣点。

用法：
  python3 scripts/gen_conversion_repurchase_xlsx.py [--out PATH]
  python3 scripts/gen_conversion_repurchase_xlsx.py --install-deps   # 自动 pip 安装 openpyxl
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

OUTPUT_PREFIX = "转化与复购运营模板"


def _ensure_openpyxl(install_deps: bool):
    """检测 openpyxl；缺失时引导安装。返回 openpyxl 模块或退出。"""
    try:
        import openpyxl  # noqa: F401
        return openpyxl
    except ImportError:
        sys.stderr.write(
            "[gen_conversion_repurchase_xlsx] ⚠️ 缺少依赖 openpyxl，无法生成 Excel。\n"
            "请先安装：pip3 install openpyxl（或 python3 -m pip install openpyxl）\n"
            "也可重跑本脚本并带 --install-deps 自动安装：\n"
            "  python3 scripts/gen_conversion_repurchase_xlsx.py --install-deps\n"
        )
        if install_deps:
            sys.stderr.write("[gen_conversion_repurchase_xlsx] 正在自动安装 openpyxl …\n")
            import subprocess
            rc = subprocess.call([sys.executable, "-m", "pip", "install", "openpyxl"])
            if rc != 0:
                sys.stderr.write("[gen_conversion_repurchase_xlsx] 自动安装失败，请手动安装后重试。\n")
                sys.exit(2)
            try:
                import openpyxl  # noqa: F401
                return openpyxl
            except ImportError:
                sys.stderr.write("[gen_conversion_repurchase_xlsx] 安装后仍导入失败，请检查 pip 环境。\n")
                sys.exit(2)
        sys.exit(2)


# 样式 ------------------------------------------------------------------
def _header_style(openpyxl):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    thin = Side(style="thin", color="BFBFBF")
    return {
        "font": Font(bold=True, color="FFFFFF", size=11),
        "fill": PatternFill("solid", fgColor="305496"),
        "align": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
    }


def _cell_style(openpyxl):
    from openpyxl.styles import Alignment, Border, Side
    thin = Side(style="thin", color="D9D9D9")
    return {
        "align": Alignment(vertical="top", wrap_text=True),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
    }


def _apply_header(ws, row_idx, n_cols, openpyxl):
    st = _header_style(openpyxl)
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = st["font"]
        cell.fill = st["fill"]
        cell.alignment = st["align"]
        cell.border = st["border"]


def _apply_body(ws, row_idx, n_cols, openpyxl):
    st = _cell_style(openpyxl)
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.alignment = st["align"]
        cell.border = st["border"]


def _write_sheet(ws, headers, samples, widths, note, openpyxl, extra_rows=18):
    from openpyxl.utils import get_column_letter
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    _apply_header(ws, 1, len(headers), openpyxl)
    for r, row in enumerate(samples, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
        _apply_body(ws, r, len(headers), openpyxl)
    for r in range(2 + len(samples), 2 + len(samples) + extra_rows):
        _apply_body(ws, r, len(headers), openpyxl)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 40
    ws.freeze_panes = "A2"
    if note:
        note_row = 2 + len(samples) + extra_rows + 1
        ws.cell(row=note_row, column=1, value=note)
        ws.cell(row=note_row, column=1).font = openpyxl.styles.Font(italic=True, color="7F7F7F")


# 各 Sheet --------------------------------------------------------------
def _sheet_readme(wb, openpyxl):
    ws = wb.create_sheet("使用说明", 0)
    lines = [
        ["转化与复购运营模板 —— 使用说明", ""],
        ["", ""],
        ["本文件由「转化与复购运营」技能生成，是一份拿来即填的空白模板。", ""],
        ["共 4 个业务 Sheet（四件套）：", ""],
        ["  1. 首购转化漏斗", "场景 × 五要素(人群包/触发条件/话术/权益/导购动作)；把你的首购场景逐行填（示例行可删）"],
        ["  2. 复购策略表", "复购品类 × 补货周期/触发条件/推荐机制/权益梯度/话术组/渠道/频次"],
        ["  3. 企微话术组", "话术ID × 五段式(开场/价值/稀缺/促销/承接)/触发/人设/答疑FAQ/二次跟进/导购动作"],
        ["  4. 沉睡召回方案", "召回次序 × 6层分级/触达节奏/权益梯度/话术组/止召回/效果归因"],
        ["", ""],
        ["填表要点（五道红线必过）：", ""],
        ["  · 品类口径先定：高频快消/中频周期/低频高客单/周期刚需，四件套默认档依赖此", ""],
        ["  · 五要素齐备：每条漏斗必含人群包+触发条件+话术+权益+导购动作，缺任一标【缺陷—需补】", ""],
        ["  · 频次红线：1V1 一周≤2次、同一客户月≤4条群发（硬限），越界标【越界—减触达】", ""],
        ["  · 召回权益≤新客70%：沉睡召回权益不超新用户注册权益70%，防'沉睡-领券-离开'恶性循环", ""],
        ["  · 沉睡分层≥3层：按'购买频次×沉睡时长'分（轻60天/中90天/深180天…），只分两层召回率<3%", ""],
        ["  · 归因用30日：效果归因用30日贡献价值，非24小时唤醒率（后者会算虚）", ""],
        ["  · 复购口径明示：复购率必标用户维度(复购用户÷总购买用户)还是订单维度(重复购买次数÷总购买次数)", ""],
        ["  · 低频高客单(车/珠宝/家居)：复购漏斗降权，高客单漏斗转向关联消费(保养/配件)+转介绍，忌打折透支品牌", ""],
        ["  · 正交引用：RFM结果引用父技能【私域用户运营】、生命周期阶段引用【用户生命周期旅程】，本技能不重算", ""],
        ["", ""],
        ["生成脚本：scripts/gen_conversion_repurchase_xlsx.py（依赖 openpyxl）", ""],
    ]
    for r, (a, b) in enumerate(lines, 1):
        ws.cell(row=r, column=1, value=a)
        ws.cell(row=r, column=2, value=b)
    ws.cell(row=1, column=1).font = openpyxl.styles.Font(bold=True, size=14, color="305496")
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 60
    ws.sheet_view.showGridLines = False


def _sheet_first_purchase(wb, openpyxl):
    ws = wb.create_sheet("首购转化漏斗")
    headers = ["场景", "人群包", "触发条件", "话术要点", "权益", "导购动作"]
    samples = [
        ["加好友破冰", "新粉", "加好友即时", "开场:礼貌询问+自我介绍+行为回顾;价值:新人福利", "新人首单券(满99-30)", "打标新粉→分流专属顾问"],
        ["首单培育", "加好友未首购D+1", "T+1", "价值点:产品价值+短期利益;稀缺:限时", "限时新人礼包(小样+券)", "推爆款商品"],
        ["首单冲刺", "加好友48h未购", "T+2/T+3", "促单:行动指令;承接:答疑FAQ", "倒计时加码券(满129-40)", "人工1v1承接答异议"],
        ["未购流失预警", "加好友7天未购", "T+7", "关怀:询问未购原因+新利益点", "无门槛小样包", "登记小结→续跟"],
    ]
    note = "说明：每行五要素齐备（缺任一标【缺陷—需补】）；新粉48h多波段促首单忌深度推荐伤体验；1V1≤2次/周；T+日序列递进每波段换素材。"
    _write_sheet(ws, headers, samples, [16, 18, 16, 40, 22, 26], note, openpyxl, extra_rows=18)


def _sheet_repurchase(wb, openpyxl):
    ws = wb.create_sheet("复购策略表")
    headers = ["复购品类", "补货周期", "复购触发条件", "推荐机制", "权益梯度", "关联话术组", "触达渠道", "频次限制"]
    samples = [
        ["护肤线(粉底/精华)", "60-90天", "周期结束前3-7天", "关联:买粉底推同系列精华", "复购券→2倍积分→会员升级奖励", "T+1复购福利/T+7续购/T+21复购/T+24倒计时", "企微1v1/社群/小程序", "1v1≤2次/周"],
        ["配件(美妆蛋/化妆刷)", "3-6月", "主品购后3月", "连带:配件短周期复购", "配件专享券(8折)", "单次复购提醒", "企微1v1", "月1次"],
        ["周期刚需(卸妆/洁面)", "30-45天", "周期结束前5天", "补货提醒+1件多瓶优惠", "买2送1装", "T+补货提醒", "企微1v1/小程序", "月1次"],
        ["低频高客单-关联复购", "如保养12月", "保养到期前3月", "关联消费:保养/配件", "服务券非打折", "服务提醒T+序列", "专属顾问1v1", "季度1次"],
    ]
    note = "说明：复购率口径必明示（用户维度=复购用户÷总购买用户；订单维度=重复购买次数÷总购买次数，两套同列）；按品类补货周期触发非拍天数；魔法数字第4次复购后忠诚稳定；1V1≤2次/周；低频高客单不出现'提复购率'强目标（否则标【打法错配】转向关联消费+转介绍）。"
    _write_sheet(ws, headers, samples, [22, 14, 20, 26, 28, 32, 22, 14], note, openpyxl, extra_rows=18)


def _sheet_script(wb, openpyxl):
    ws = wb.create_sheet("企微话术组")
    headers = ["话术ID", "适用人群包", "触发场景", "人设", "发送渠道", "发送时段", "开场白", "价值点", "稀缺/促单", "承接答疑FAQ", "是否二次跟进", "关联权益", "导购动作"]
    samples = [
        ["HS-REPURCHASE-D1", "复购2次+护肤线", "T+1(补货周期前5天)", "福利官", "企微1v1", "晚8点(直播前)", "Hi~上次买的XX用得怎么样，该补货啦", "短期:复购立减30+2倍积分；长期:新品内测", "活动倒计时3天，前100名加赠小样", "价格异议/效果疑问/拒绝挽回2轮", "二轮(间隔3天)", "复购券(满199-30)+2倍积分", "打标复购意向→推组货→登记小结"],
        ["HS-NEWUSER-D0", "新粉", "加好友即时", "福利官", "企微1v1", "任何时段", "Hi~我是专属顾问，欢迎领取新人福利", "新人专享福利+长期服务", "新人限时礼包", "色号/肤质疑问/拒绝挽回2轮", "一轮", "新人首单券", "打标新粉→分流专属顾问"],
        ["HS-RECALL-T1", "轻度沉睡60天", "T+1关怀", "专属顾问", "企微1v1", "晚8点", "Hi~好久不见，给你留了专属福利", "新品福利+专属权益", "3天时限", "顾虑/拒绝挽回2轮", "二轮(间隔T+7)", "关怀券(满99-15)", "打标召回意向→推回归礼→登记小结"],
    ]
    note = "说明：五段式齐全（开场/价值/稀缺/促销/承接）；答疑FAQ含拒绝挽回2轮；1V1≤2次/周；导购动作必含'登记小结'闭环（华为云：必登记才算任务完成）；'是否二次跟进'填一轮/二轮/三轮+间隔。"
    _write_sheet(ws, headers, samples, [20, 20, 20, 12, 14, 16, 30, 30, 28, 24, 16, 26, 28], note, openpyxl, extra_rows=18)


def _sheet_recall(wb, openpyxl):
    ws = wb.create_sheet("沉睡召回方案")
    headers = ["召回次序", "人群分层", "触达节奏", "触达渠道", "权益梯度", "话术组", "止召回条件", "效果指标"]
    samples = [
        ["第1次", "轻度沉睡(60天)", "T+1关怀", "企微1v1", "关怀券(轻量,满99-15)", "唤醒4要素(理由+权益+行动+时限)", "已复购/明确拒绝", "30日复购率"],
        ["第2次", "第1次未响应", "T+7(≥24h)", "短信", "积分加价购(低门槛换购)", "新利益点", "3次无响应", "30日贡献价值"],
        ["第3次", "中度沉睡(90天)", "T+15(≥24h)", "AI外呼", "召回券(≤新客70%)", "专属召回+情感", "拉黑", "7日留存率"],
        ["第4次", "深度沉睡(180天,仅高价值)", "约7天后", "人工1v1", "历史偏好召回礼", "专属方案", "-", "召回ROI"],
        ["放弃", "幽灵用户(注册未购)/流失高危", "-", "-", "不触达省成本", "-", "直接放弃", "-"],
    ]
    note = "说明：沉睡定义与类目补货周期相关不可一刀切；召回分层≥3层（6层：轻60/中90/深180/极深180+/幽灵/流失高危）；权益≤新客70%硬约束；归因用30日贡献价值非24小时唤醒率；前3次每次≥24h间隔；放弃幽灵用户省40%触达成本。"
    _write_sheet(ws, headers, samples, [12, 26, 16, 16, 26, 30, 18, 18], note, openpyxl, extra_rows=18)


def build_workbook(openpyxl):
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)
    _sheet_readme(wb, openpyxl)
    _sheet_first_purchase(wb, openpyxl)
    _sheet_repurchase(wb, openpyxl)
    _sheet_script(wb, openpyxl)
    _sheet_recall(wb, openpyxl)
    return wb


def write_to_disk(wb, out_path):
    """三级兜底：--out → 当前目录 → /tmp。返回绝对路径或空。"""
    candidates = []
    if out_path:
        p = Path(out_path).expanduser()
        if p.is_dir():
            p = p / (OUTPUT_PREFIX + ".xlsx")
        candidates.append(p)
    candidates.append(Path.cwd() / (OUTPUT_PREFIX + ".xlsx"))
    candidates.append(Path("/tmp") / (OUTPUT_PREFIX + ".xlsx"))
    last_err = None
    for p in candidates:
        try:
            p.parent.mkdir(parents=True, exist_ok=True)
            wb.save(str(p))
            return str(p.resolve())
        except Exception as e:  # noqa: BLE001
            last_err = e
            continue
    if last_err:
        sys.stderr.write(f"[gen_conversion_repurchase_xlsx] 写盘失败：{last_err}\n")
    return ""


def main():
    ap = argparse.ArgumentParser(description="生成「转化与复购运营」Excel 空白模板")
    ap.add_argument("--out", default=None, help="输出 xlsx 路径（目录或文件均可，可省略走兜底）")
    ap.add_argument("--install-deps", action="store_true", help="缺失 openpyxl 时自动 pip 安装")
    args = ap.parse_args()

    openpyxl = _ensure_openpyxl(args.install_deps)
    wb = build_workbook(openpyxl)
    out = write_to_disk(wb, args.out)
    if out:
        print(f"CONVERSION_REPURCHASE_XLSX_FILE={out}")
        sys.exit(0)
    sys.exit(1)


if __name__ == "__main__":
    main()
