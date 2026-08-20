#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
公私域联动方案生成 · 四件套 Excel 空白模板生成器
照搬技能四件套 Markdown 结构成 4 个 Sheet：触点地图 / 承接链路 / 渠道优先级RICE / 承接期SOP。

纯本地、免 Key。依赖 openpyxl；缺依赖时退出码 2 并给安装引导，--install-deps 可自动装。
三级兜底写盘：--out 指定路径 → 当前工作目录 → /tmp。成功退出码 0，
stdout 协议行：TOUCHPOINT_XLSX_FILE=<绝对路径> 供 agent 解析。
"""

import sys
import os
import subprocess

XLSX_FILENAME = "公私域联动方案生成四件套模板.xlsx"

# ---------- 依赖检测 ----------
def ensure_openpyxl(auto_install=False):
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        sys.stderr.write(
            "\n[缺依赖] 本脚本需要 openpyxl 来生成 Excel。\n"
            "安装方法（任选其一）：\n"
            "  1) python3 -m pip install openpyxl\n"
            "  2) python3 {script} --install-deps   # 本脚本自动安装\n"
            "装好后重新运行本脚本即可。\n\n".format(script=sys.argv[0])
        )
        if auto_install:
            sys.stderr.write("[自动安装] 正在执行: python3 -m pip install openpyxl ...\n")
            try:
                subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
                return True
            except Exception as e:  # noqa: BLE001
                sys.stderr.write("[自动安装失败] {}\n请手动安装：python3 -m pip install openpyxl\n".format(e))
        sys.exit(2)


# ---------- 写盘路径解析（三级兜底）----------
def resolve_out_path(out_arg):
    if out_arg:
        out_dir = os.path.dirname(os.path.abspath(out_arg)) or "."
        if os.access(out_dir if os.path.isdir(out_dir) else os.path.dirname(out_dir) or os.getcwd(), os.W_OK):
            return os.path.abspath(out_arg)
        # 指定路径不可写，降级
    # 兜底 1：当前工作目录
    try:
        cwd = os.getcwd()
        test = os.path.join(cwd, ".wb_touchpoint_writetest")
        with open(test, "w") as f:
            f.write("x")
        os.remove(test)
        return os.path.join(cwd, XLSX_FILENAME)
    except OSError:
        pass
    # 兜底 2：/tmp
    return os.path.join("/tmp", XLSX_FILENAME)


# ---------- Sheet 定义 ----------
SHEET_TOUCHPOINT_HEADERS = [
    "#", "触点名称", "所属平台载体", "触点类型", "生命周期阶段",
    "承接动作(用户侧+运营侧)", "沉淀位置", "参考指标", "负责人",
    "优先级", "触达频次上限", "引流诱饵", "转化目标KPI", "关联触点",
]
SHEET_TOUCHPOINT_NOTES = [
    "触点类型：公域 / 半私域 / 私域 / 线下",
    "生命周期阶段：引入期 / 成长期 / 成熟期 / 休眠期 / 流失期",
    "参考指标用区间，如 公转私沉淀5-15%、支付后关注20-35%（标【按业务校准】）",
    "质量检查：核心触点四类每类≥1个，缺整类追问",
]

SHEET_LINK_HEADERS = ["阶段", "节点(触点)", "用户动作", "承接动作", "沉淀位置", "后续触达"]
SHEET_LINK_ROWS = [
    ["1.公域种草", "", "", "", "", ""],
    ["2.私域沉淀", "", "", "", "", ""],
    ["3.活跃培育", "", "", "", "", ""],
    ["4.转化变现", "", "", "", "", ""],
    ["5.会员成长", "", "", "", "", ""],
    ["6.流失召回", "", "", "", "", ""],
]
SHEET_LINK_NOTES = [
    "6阶段每阶段≥1个触点节点",
    "质量检查：每节点『沉淀位置』=下阶段『入口』（无断点）；断点标红补承接动作",
    "边界：会员成长转会员积分技能；流失召回转分层触达技能",
]

SHEET_RICE_HEADERS = [
    "渠道", "类型", "Reach(规模×漏斗)", "Impact(3/2/1/0.5/0.25)",
    "Confidence(100/80/50/20%)", "Effort(人月)", "Score=(R×I×C)÷E", "优先级档",
]
SHEET_RICE_NOTES = [
    "Reach=可承载规模×转化漏斗（视频号Reach高但转企微5-15%；企微Reach低但首购20-35%）",
    "Impact：3=直接影响付费/留存 / 2=中 / 1=轻 / 0.5·0.25=极低",
    "Confidence：100%有数据 / 80%中等 / 50%估算 / 20%=登月",
    "Effort：人月（研发+设计+运营），最小0.5",
    "优先级档：必做(Score前2) / 优先 / 观察 / 暂缓(Score末2)",
    "质量检查：R/I/C/E均给数值，禁『高/中/低』模糊词；渠道类型按业务反推（高频低客单→社群1v多）",
]

SHEET_SOP_HEADERS = [
    "SOP名称", "适用阶段", "触发条件(用户状态事件)", "触达渠道", "触达时间间隔",
    "文案模板(≤100字)", "内容类型", "负责人", "数据指标", "异常处理",
    "复盘节点", "版本号",
]
SHEET_SOP_SEED_ROWS = [
    ["新客7天承接SOP", "新客承接", "加好友通过时（非定时）", "企微1v1", "Day1/3/7",
     "Day1欢迎+新人券 / Day3品牌故事 / Day7首单钩子", "品牌背书/产品教育/利益驱动",
     "社群运营", "打开率/首购率", "回复→暂停SOP;咨询→转人工", "每周", "V1.0"],
    ["首购转化SOP", "首购转化", "浏览未购3天/首购后第3天", "企微1v1", "相对触发",
     "浏览未购『你关注的XX今天半价』", "利益驱动", "社群运营", "首购率/复购率",
     "同上", "每周", "V1.0"],
    ["跨触点编排SOP", "跨触点编排", "公众号关注后/社群活跃后", "企微+社群+小程序", "状态流转",
     "按触点状态推送对应承接话术", "品牌背书/产品教育", "运营", "跨触点转化率",
     "同上", "每月", "V1.0"],
]
SHEET_SOP_NOTES = [
    "承接期聚焦：新客7天承接 / 首购转化 / 跨触点编排三主干",
    "质量检查：触发条件必须是用户状态事件，非定时（每天8点发=反模式→异常5）",
    "边界：成熟期/沉睡召回全生命周期分层触达归分层触达技能，不写进本表",
    "文案≤100字，一条一核心点；结构=钩子+价值点+CTA+素材",
    "内容比例参考 营销:互动:关怀:日常=3:3:3:1",
]


def _set_col_widths(ws, widths):
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _write_notes(ws, notes, start_row):
    for i, note in enumerate(notes):
        ws.cell(row=start_row + i, column=1, value="📝 " + note)


def build_touchpoint_sheet(wb):
    from openpyxl.styles import Font, PatternFill, Alignment
    ws = wb.create_sheet("1-触点地图")
    ws.append(["公私域联动方案生成 · 触点地图（13列盘点表）"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["触点全枚举参考：公域(朋友圈广告/视频号原生广告/搜一搜/抖音小红书种草/电商包裹) · "
               "半私域(公众号菜单推文模板消息/视频号主页直播间私信) · "
               "私域(企微1v1朋友圈社群活码/导购个微/小程序首页详情页支付页订阅消息/会员系统) · "
               "线下(门店收银扫码/POP导购码/包裹卡/活动)"])
    ws.append([])
    ws.append(SHEET_TOUCHPOINT_HEADERS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for _ in range(15):
        ws.append([""] * len(SHEET_TOUCHPOINT_HEADERS))
    _write_notes(ws, SHEET_TOUCHPOINT_NOTES, ws.max_row + 2)
    _set_col_widths(ws, [4, 18, 16, 12, 12, 28, 14, 18, 10, 10, 14, 14, 14, 16])


def build_link_sheet(wb):
    from openpyxl.styles import Font, PatternFill, Alignment
    ws = wb.create_sheet("2-承接链路图")
    ws.append(["公私域联动方案生成 · 承接链路图（6阶段）"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(SHEET_LINK_HEADERS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in SHEET_LINK_ROWS:
        ws.append(row)
    _write_notes(ws, SHEET_LINK_NOTES, ws.max_row + 2)
    _set_col_widths(ws, [14, 22, 18, 26, 20, 22])


def build_rice_sheet(wb):
    from openpyxl.styles import Font, PatternFill, Alignment
    ws = wb.create_sheet("3-渠道优先级RICE")
    ws.append(["公私域联动方案生成 · 渠道优先级（RICE评分表）"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["公式：Score=(Reach×Impact×Confidence)÷Effort，从大到小排序"])
    ws.append([])
    ws.append(SHEET_RICE_HEADERS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for _ in range(10):
        ws.append([""] * len(SHEET_RICE_HEADERS))
    _write_notes(ws, SHEET_RICE_NOTES, ws.max_row + 2)
    _set_col_widths(ws, [16, 10, 24, 18, 18, 14, 16, 14])


def build_sop_sheet(wb):
    from openpyxl.styles import Font, PatternFill, Alignment
    ws = wb.create_sheet("4-承接期SOP")
    ws.append(["公私域联动方案生成 · 承接期触达SOP（12字段表）"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(SHEET_SOP_HEADERS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in SHEET_SOP_SEED_ROWS:
        ws.append(row)
    for _ in range(8):
        ws.append([""] * len(SHEET_SOP_HEADERS))
    _write_notes(ws, SHEET_SOP_NOTES, ws.max_row + 2)
    _set_col_widths(ws, [16, 12, 22, 14, 14, 28, 18, 10, 16, 22, 10, 8])


def main():
    args = sys.argv[1:]
    auto_install = "--install-deps" in args
    out_arg = None
    for a in args:
        if a.startswith("--out="):
            out_arg = a.split("=", 1)[1]
        elif a == "--out" and args.index(a) + 1 < len(args):
            out_arg = args[args.index(a) + 1]

    if "--help" in args or "-h" in args:
        print("用法：python3 gen_touchpoint_xlsx.py [--out=路径] [--install-deps]")
        print("生成公私域联动方案生成四件套 Excel 空白模板（4 Sheet）。")
        sys.exit(0)

    if not ensure_openpyxl(auto_install=auto_install):
        sys.exit(2)
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # 删除默认 Sheet
    build_touchpoint_sheet(wb)
    build_link_sheet(wb)
    build_rice_sheet(wb)
    build_sop_sheet(wb)

    out_path = resolve_out_path(out_arg)
    try:
        wb.save(out_path)
    except OSError:
        # 写盘失败，最后兜底 /tmp（保证总能生成）
        out_path = os.path.join("/tmp", XLSX_FILENAME)
        wb.save(out_path)

    print("TOUCHPOINT_XLSX_FILE={}".format(os.path.abspath(out_path)))
    print("已生成公私域联动方案生成四件套空白模板（4 Sheet）：")
    print("  1-触点地图 / 2-承接链路图 / 3-渠道优先级RICE / 4-承接期SOP")
    sys.exit(0)


if __name__ == "__main__":
    main()
