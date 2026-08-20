#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_action_tracker.py — 生成本周追踪事项 Excel。

与 weekly_minutes_docx.py（或 meeting_docx.py）共用 content JSON，仅读取 action_items。
状态列支持下拉；选“完成”时，该条整行自动填充浅绿色。

用法：
  python build_action_tracker.py --content content.json --out 甲方A-本周追踪事项-20260801.xlsx
"""

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


SHEET_NAME = "本周追踪事项"
HEADERS = ["编号", "会议安排事项", "负责人", "追踪人", "完成时限", "状态", "来源", "备注"]
STATUS_OPTIONS = ["未启动", "进行中", "阻塞", "完成", "取消"]
STATUS_COLUMN = 6
MAX_TRACKER_ROW = 1000
COMPLETED_FILL = "FFE2F0D9"


def responsible(item):
    """优先读取“负责人”，兼容旧 JSON 的“责任人”键。"""
    return str(item.get("负责人", item.get("责任人", ""))).strip()


def normalize_status(value):
    status = str(value or "未启动").strip()
    if status == "已完成":
        status = "完成"
    if status not in STATUS_OPTIONS:
        raise ValueError(
            f"状态 {status!r} 无效，必须是：" + "/".join(STATUS_OPTIONS)
        )
    return status


def normalize_items(content):
    items = content.get("action_items", [])
    if not isinstance(items, list):
        raise ValueError("action_items 必须是列表")
    if not items:
        raise ValueError("本周待办为空，不能生成本周追踪事项 Excel")

    normalized = []
    seen_ids = set()
    for index, item in enumerate(items, 1):
        if not isinstance(item, dict):
            raise ValueError(f"第 {index} 条 action_items 必须是对象")
        item_id = str(item.get("编号", f"本周-{index:02d}")).strip()
        subject = str(item.get("事项", "")).strip()
        owner = responsible(item)
        deadline = str(item.get("完成时限", "")).strip()
        if not item_id or not subject or not owner or not deadline:
            raise ValueError(f"第 {index} 条缺少编号、事项、负责人或完成时限")
        if item_id in seen_ids:
            raise ValueError(f"本周待办编号重复：{item_id}")
        seen_ids.add(item_id)
        normalized.append(
            [
                item_id,
                subject,
                owner,
                str(item.get("追踪人", "")).strip(),
                deadline,
                normalize_status(item.get("状态")),
                str(item.get("来源", "")).strip(),
                str(item.get("备注", "")).strip(),
            ]
        )
    return normalized


def build_tracker(content, output_path):
    rows = normalize_items(content)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=10.5, color="000000")
    thin_gray = Side(style="thin", color="B7C9D6")
    border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.cell(1, STATUS_COLUMN).comment = Comment(
        "请从下拉菜单选择状态；选“完成”后整行自动变为浅绿色。",
        "Moways",
    )

    for row_index, values in enumerate(rows, 2):
        for col_index, value in enumerate(values, 1):
            cell = ws.cell(row=row_index, column=col_index, value=value)
            cell.font = body_font
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center" if col_index in {1, 3, 4, 5, 6} else "left",
                vertical="top",
                wrap_text=True,
            )
        ws.row_dimensions[row_index].height = 32

    ws.row_dimensions[1].height = 28
    widths = [12, 42, 13, 13, 18, 12, 18, 30]
    for col_index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col_index)].width = width

    last_data_row = 1 + len(rows)
    table = Table(displayName="WeeklyActionTracker", ref=f"A1:H{last_data_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(STATUS_OPTIONS) + '"',
        allow_blank=False,
    )
    validation.error = "请从下拉列表选择状态。"
    validation.errorTitle = "状态值无效"
    validation.prompt = "可选：未启动、进行中、阻塞、完成、取消。"
    validation.promptTitle = "更新事项状态"
    validation.showErrorMessage = True
    validation.showInputMessage = True
    ws.add_data_validation(validation)
    validation.add(f"F2:F{MAX_TRACKER_ROW}")

    completed_fill = PatternFill("solid", fgColor=COMPLETED_FILL, bgColor=COMPLETED_FILL)
    completed_rule = FormulaRule(formula=['$F2="完成"'], fill=completed_fill)
    ws.conditional_formatting.add(f"A2:H{MAX_TRACKER_ROW}", completed_rule)

    ws.auto_filter.ref = f"A1:H{last_data_row}"
    ws.print_title_rows = "1:1"
    ws.print_area = f"A1:H{last_data_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "第 &[Page] 页，共 &[Pages] 页"

    wb.properties.creator = "Moways"
    wb.properties.title = str(content.get("title", "本周追踪事项")) + " · 追踪事项"
    wb.save(output_path)
    return output_path


def validate_workbook(path, expected_rows):
    """重新打开输出并检查结构，避免交互规则在保存后丢失。"""
    wb = load_workbook(path)
    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(f"缺少工作表：{SHEET_NAME}")
    ws = wb[SHEET_NAME]
    if [ws.cell(1, col).value for col in range(1, 9)] != HEADERS:
        raise ValueError("Excel 表头不符合规格")
    if ws.max_row - 1 != expected_rows:
        raise ValueError("Excel 数据行数与 action_items 不一致")
    if not ws.tables:
        raise ValueError("Excel 未建立可筛选表格")
    if len(ws.data_validations.dataValidation) != 1:
        raise ValueError("Excel 状态下拉规则缺失")
    conditional_ranges = list(ws.conditional_formatting)
    if not conditional_ranges:
        raise ValueError("Excel 完成行浅绿条件格式缺失")
    rules = ws.conditional_formatting[conditional_ranges[0]]
    if rules[0].formula != ['$F2="完成"']:
        raise ValueError("Excel 完成行条件公式错误")
    if rules[0].dxf.fill.fgColor.rgb != COMPLETED_FILL:
        raise ValueError("Excel 完成行浅绿颜色或透明度错误")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--content", required=True, help="与纪要 docx 生成共用的内容 JSON")
    parser.add_argument("--out", required=True, help="输出 xlsx 路径")
    args = parser.parse_args()

    with open(args.content, encoding="utf-8-sig") as handle:
        content = json.load(handle)
    rows = normalize_items(content)
    output = build_tracker(content, args.out)
    validate_workbook(output, len(rows))
    print(output)


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"生成本周追踪事项 Excel 失败：{exc}", file=sys.stderr)
        raise
