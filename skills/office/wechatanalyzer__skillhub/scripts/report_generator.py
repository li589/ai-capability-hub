#!/usr/bin/env python3
"""
报告生成器 - 生成HTML、PPTX、Word格式的分析报告
"""

import json
import os
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, Optional

from jinja2 import Template


class ReportGenerator:
    """分析报告生成器"""

    def __init__(self, config: Dict[str, Any]):
        self.config = config
        self.template_dir = Path(__file__).parent.parent / "templates"
        self.output_dir = Path(config.get('report', {}).get('output_dir', 'data/reports'))
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def generate(self, results: Dict[str, Any], report_type: str = 'html',
                 period: str = 'weekly') -> str:
        """生成报告"""

        if report_type == 'html':
            return self._generate_html(results, period)
        elif report_type == 'pptx':
            return self._generate_pptx(results, period)
        elif report_type == 'word':
            return self._generate_word(results, period)
        elif report_type == 'all':
            outputs = []
            outputs.append(self._generate_html(results, period))
            outputs.append(self._generate_pptx(results, period))
            outputs.append(self._generate_word(results, period))
            return '; '.join(outputs)

        return ""

    def _generate_html(self, results: Dict[str, Any], period: str) -> str:
        """生成HTML报告"""
        template_path = self.template_dir / "report_template.html"

        if template_path.exists():
            with open(template_path, 'r', encoding='utf-8') as f:
                template_content = f.read()
        else:
            template_content = self._get_default_html_template()

        template = Template(template_content)

        # 准备数据
        stats = results.get('stats', {})
        mbti = results.get('mbti', {})
        big_five = results.get('big_five', {})
        sentiment = results.get('sentiment', {})
        prediction = results.get('prediction', {})

        context = {
            'title': f'微信聊天分析报告',
            'generated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
            'period': self._get_period_label(period),
            'stats': stats,
            'mbti': mbti,
            'big_five': big_five,
            'sentiment': sentiment,
            'prediction': prediction,
            'time_distribution': results.get('time_distribution', {}),
            'word_frequency': results.get('word_frequency', [])[:20]
        }

        html_content = template.render(**context)

        # 保存文件
        output_path = self.output_dir / f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html_content)

        return str(output_path)

    def _generate_pptx(self, results: Dict[str, Any], period: str) -> str:
        """生成PPTX报告"""
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RgbColor
        except ImportError:
            return "错误: python-pptx未安装，请使用 pip install python-pptx"

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        # 标题页
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        title = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(2))
        title.text = "微信聊天分析报告"
        title.text_frame.paragraphs[0].font.size = Pt(44)
        title.text_frame.paragraphs[0].font.bold = True

        subtitle = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11), Inches(1))
        subtitle.text = f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n周期: {self._get_period_label(period)}"
        subtitle.text_frame.paragraphs[0].font.size = Pt(18)

        # 统计页
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        stats = results.get('stats', {})
        self._add_slide_title(slide, "聊天统计")
        stats_text = f"""
总消息数: {stats.get('total_messages', 0)}
你的消息: {stats.get('self_messages', 0)}
对方消息: {stats.get('other_messages', 0)}
平均消息长度: {stats.get('avg_message_length', 0)} 字符
对话持续时间: {stats.get('conversation_duration_days', 0)} 天
"""
        self._add_text_box(slide, stats_text, Inches(1), Inches(2), Inches(11), Inches(5))

        # MBTI页
        if results.get('mbti'):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self._add_slide_title(slide, "MBTI人格分析")
            mbti = results.get('mbti', {})
            mbti_text = f"""
人格类型: {mbti.get('type', '未知')}
类型名称: {mbti.get('name', '未知')}
置信度: {mbti.get('confidence', 0)}%
描述: {mbti.get('description', '')}
"""
            self._add_text_box(slide, mbti_text, Inches(1), Inches(2), Inches(11), Inches(5))

        # 大五人格页
        if results.get('big_five'):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self._add_slide_title(slide, "大五人格分析")
            bf = results.get('big_five', {})
            bf_text = f"""
开放性: {bf.get('openness', 50)}%
尽责性: {bf.get('conscientiousness', 50)}%
外向性: {bf.get('extraversion', 50)}%
宜人性: {bf.get('agreeableness', 50)}%
神经质: {bf.get('neuroticism', 50)}%
"""
            self._add_text_box(slide, bf_text, Inches(1), Inches(2), Inches(11), Inches(5))

        # 对话预测页
        if results.get('prediction'):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            self._add_slide_title(slide, "对话预测")
            pred = results.get('prediction', {})
            pred_text = f"对方可能会说:\n{pred.get('next_message', '分析中...')}"
            self._add_text_box(slide, pred_text, Inches(1), Inches(2), Inches(11), Inches(5))

        # 保存
        output_path = self.output_dir / f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pptx"
        prs.save(str(output_path))
        return str(output_path)

    def _generate_word(self, results: Dict[str, Any], period: str) -> str:
        """生成Word报告"""
        try:
            from docx import Document
            from docx.shared import Inches, Pt
        except ImportError:
            return "错误: python-docx未安装，请使用 pip install python-docx"

        doc = Document()

        # 标题
        doc.add_heading('微信聊天分析报告', 0)

        # 基本信息
        doc.add_heading('基本信息', level=1)
        doc.add_paragraph(f'生成时间: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
        doc.add_paragraph(f'报告周期: {self._get_period_label(period)}')

        # 聊天统计
        doc.add_heading('聊天统计', level=1)
        stats = results.get('stats', {})
        doc.add_paragraph(f"总消息数: {stats.get('total_messages', 0)}")
        doc.add_paragraph(f"你的消息: {stats.get('self_messages', 0)}")
        doc.add_paragraph(f"对方消息: {stats.get('other_messages', 0)}")
        doc.add_paragraph(f"平均消息长度: {stats.get('avg_message_length', 0)} 字符")
        doc.add_paragraph(f"对话持续时间: {stats.get('conversation_duration_days', 0)} 天")

        # MBTI分析
        if results.get('mbti'):
            doc.add_heading('MBTI人格分析', level=1)
            mbti = results.get('mbti', {})
            doc.add_paragraph(f"人格类型: {mbti.get('type', '未知')}")
            doc.add_paragraph(f"类型名称: {mbti.get('name', '未知')}")
            doc.add_paragraph(f"置信度: {mbti.get('confidence', 0)}%")
            doc.add_paragraph(f"描述: {mbti.get('description', '')}")

        # 大五人格分析
        if results.get('big_five'):
            doc.add_heading('大五人格分析', level=1)
            bf = results.get('big_five', {})
            doc.add_paragraph(f"开放性: {bf.get('openness', 50)}%")
            doc.add_paragraph(f"尽责性: {bf.get('conscientiousness', 50)}%")
            doc.add_paragraph(f"外向性: {bf.get('extraversion', 50)}%")
            doc.add_paragraph(f"宜人性: {bf.get('agreeableness', 50)}%")
            doc.add_paragraph(f"神经质: {bf.get('neuroticism', 50)}%")

        # 情感分析
        if results.get('sentiment'):
            doc.add_heading('情感分析', level=1)
            sent = results.get('sentiment', {})
            doc.add_paragraph(f"正面情感: {sent.get('positive', 50)}%")
            doc.add_paragraph(f"负面情感: {sent.get('negative', 50)}%")
            doc.add_paragraph(f"情感趋势: {sent.get('trend', '稳定')}")

        # 对话预测
        if results.get('prediction'):
            doc.add_heading('对话预测', level=1)
            pred = results.get('prediction', {})
            doc.add_paragraph(f"对方可能会说: {pred.get('next_message', '分析中...')}")

        # 保存
        output_path = self.output_dir / f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        doc.save(str(output_path))
        return str(output_path)

    def _add_slide_title(self, slide, title: str):
        """添加幻灯片标题"""
        title_shape = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(12), Inches(1))
        title_shape.text = title
        title_shape.text_frame.paragraphs[0].font.size = Pt(32)
        title_shape.text_frame.paragraphs[0].font.bold = True

    def _add_text_box(self, slide, text: str, left: float, top: float, width: float, height: float):
        """添加文本框"""
        from pptx.util import Inches, Pt
        text_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        text_box.text = text
        for paragraph in text_box.text_frame.paragraphs:
            paragraph.font.size = Pt(18)

    def _get_period_label(self, period: str) -> str:
        """获取周期标签"""
        labels = {
            'daily': '每日报告',
            'weekly': '每周报告',
            'monthly': '每月报告',
            'quarterly': '季度报告',
            'semi-annually': '半年度报告',
            'annually': '年度报告'
        }
        return labels.get(period, period)

    def _get_default_html_template(self) -> str:
        """获取默认HTML模板"""
        return '''
<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{{ title }}</title>
    <script src="https://cdn.jsdelivr.net/npm/echarts@5.4.3/dist/echarts.min.js"></script>
    <style>
        * { margin: 0; padding: 0; box-sizing: border-box; }
        body { font-family: "Microsoft YaHei", Arial, sans-serif; background: #f5f5f5; padding: 20px; }
        .container { max-width: 1200px; margin: 0 auto; }
        .header { background: linear-gradient(135deg, #667eea 0%, #764ba2 100%); color: white; padding: 30px; border-radius: 10px; margin-bottom: 20px; }
        .header h1 { font-size: 28px; margin-bottom: 10px; }
        .header .meta { opacity: 0.9; font-size: 14px; }
        .card { background: white; border-radius: 10px; padding: 20px; margin-bottom: 20px; box-shadow: 0 2px 10px rgba(0,0,0,0.1); }
        .card h2 { color: #333; font-size: 20px; margin-bottom: 15px; border-left: 4px solid #667eea; padding-left: 10px; }
        .stats-grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(150px, 1fr)); gap: 15px; }
        .stat-item { background: #f8f9fa; padding: 15px; border-radius: 8px; text-align: center; }
        .stat-item .value { font-size: 28px; font-weight: bold; color: #667eea; }
        .stat-item .label { font-size: 12px; color: #666; margin-top: 5px; }
        .mbti-result { display: flex; align-items: center; gap: 20px; }
        .mbti-type { font-size: 48px; font-weight: bold; color: #667eea; }
        .mbti-info .name { font-size: 24px; color: #333; }
        .mbti-info .desc { color: #666; margin-top: 5px; }
        .mbti-info .confidence { color: #999; font-size: 14px; margin-top: 5px; }
        .bigfive { display: flex; gap: 20px; flex-wrap: wrap; }
        .bigfive-item { flex: 1; min-width: 120px; }
        .bigfive-item .label { font-size: 14px; color: #666; }
        .bigfive-item .bar { height: 8px; background: #e9ecef; border-radius: 4px; margin-top: 5px; }
        .bigfive-item .bar .fill { height: 100%; background: linear-gradient(90deg, #667eea, #764ba2); border-radius: 4px; }
        .bigfive-item .value { font-size: 18px; font-weight: bold; color: #333; margin-top: 5px; }
        .prediction { background: #f8f9fa; padding: 20px; border-radius: 8px; border-left: 4px solid #764ba2; }
        .prediction .quote { font-size: 18px; color: #333; font-style: italic; }
        .word-cloud { display: flex; flex-wrap: wrap; gap: 10px; }
        .word-cloud span { padding: 5px 12px; border-radius: 15px; font-size: 12px; }
        .chart { width: 100%; height: 300px; }
        .sentiment-indicator { display: flex; gap: 10px; margin-top: 10px; }
        .sentiment-indicator span { padding: 5px 15px; border-radius: 15px; font-size: 12px; }
        .positive { background: #d4edda; color: #155724; }
        .negative { background: #f8d7da; color: #721c24; }
        .neutral { background: #fff3cd; color: #856404; }
        .footer { text-align: center; color: #999; font-size: 12px; margin-top: 30px; }
    </style>
</head>
<body>
    <div class="container">
        <div class="header">
            <h1>{{ title }}</h1>
            <div class="meta">生成时间: {{ generated_at }} | {{ period }}</div>
        </div>

        <div class="card">
            <h2>聊天统计</h2>
            <div class="stats-grid">
                <div class="stat-item">
                    <div class="value">{{ stats.total_messages }}</div>
                    <div class="label">总消息数</div>
                </div>
                <div class="stat-item">
                    <div class="value">{{ stats.self_messages }}</div>
                    <div class="label">你的消息</div>
                </div>
                <div class="stat-item">
                    <div class="value">{{ stats.other_messages }}</div>
                    <div class="label">对方消息</div>
                </div>
                <div class="stat-item">
                    <div class="value">{{ stats.avg_message_length }}</div>
                    <div class="label">平均长度</div>
                </div>
                <div class="stat-item">
                    <div class="value">{{ stats.conversation_duration_days or 0 }}</div>
                    <div class="label">对话天数</div>
                </div>
            </div>
        </div>

        <div class="card">
            <h2>MBTI人格分析</h2>
            {% if mbti.type %}
            <div class="mbti-result">
                <div class="mbti-type">{{ mbti.type }}</div>
                <div class="mbti-info">
                    <div class="name">{{ mbti.name }}</div>
                    <div class="desc">{{ mbti.description }}</div>
                    <div class="confidence">置信度: {{ mbti.confidence }}%</div>
                </div>
            </div>
            {% else %}
            <p>暂无MBTI数据</p>
            {% endif %}
        </div>

        <div class="card">
            <h2>大五人格分析</h2>
            {% if big_five.openness %}
            <div class="bigfive">
                <div class="bigfive-item">
                    <div class="label">开放性</div>
                    <div class="bar"><div class="fill" style="width: {{ big_five.openness }}%"></div></div>
                    <div class="value">{{ big_five.openness }}%</div>
                </div>
                <div class="bigfive-item">
                    <div class="label">尽责性</div>
                    <div class="bar"><div class="fill" style="width: {{ big_five.conscientiousness }}%"></div></div>
                    <div class="value">{{ big_five.conscientiousness }}%</div>
                </div>
                <div class="bigfive-item">
                    <div class="label">外向性</div>
                    <div class="bar"><div class="fill" style="width: {{ big_five.extraversion }}%"></div></div>
                    <div class="value">{{ big_five.extraversion }}%</div>
                </div>
                <div class="bigfive-item">
                    <div class="label">宜人性</div>
                    <div class="bar"><div class="fill" style="width: {{ big_five.agreeableness }}%"></div></div>
                    <div class="value">{{ big_five.agreeableness }}%</div>
                </div>
                <div class="bigfive-item">
                    <div class="label">神经质</div>
                    <div class="bar"><div class="fill" style="width: {{ big_five.neuroticism }}%"></div></div>
                    <div class="value">{{ big_five.neuroticism }}%</div>
                </div>
            </div>
            {% else %}
            <p>暂无大五人格数据</p>
            {% endif %}
        </div>

        <div class="card">
            <h2>情感分析</h2>
            {% if sentiment.positive %}
            <div class="sentiment-indicator">
                <span class="positive">正面 {{ sentiment.positive }}%</span>
                <span class="negative">负面 {{ sentiment.negative }}%</span>
                <span class="neutral">中性 {{ sentiment.neutral }}%</span>
            </div>
            <p style="margin-top: 15px; color: #666;">情感趋势: {% if sentiment.trend == 'up' %}上升 ↑{% elif sentiment.trend == 'down' %}下降 ↓{% else %}稳定 →{% endif %}</p>
            {% else %}
            <p>暂无情感数据</p>
            {% endif %}
        </div>

        <div class="card">
            <h2>对话预测</h2>
            {% if prediction.next_message %}
            <div class="prediction">
                <p class="quote">对方可能会说: {{ prediction.next_message }}</p>
            </div>
            {% else %}
            <p>暂无预测数据</p>
            {% endif %}
        </div>

        <div class="card">
            <h2>高频词</h2>
            <div class="word-cloud">
                {% for word in word_frequency %}
                <span style="background: #e8f4f8; color: #2c5282;">{{ word }}</span>
                {% endfor %}
            </div>
        </div>

        <div class="footer">
            <p>由 微信聊天分析助手 生成 | 数据仅保存在本地</p>
        </div>
    </div>
</body>
</html>
        '''
