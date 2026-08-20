"""scripts.file_importer 多格式文件解析测试（v2.2.0）

覆盖 v2.2.0 新增/修复的格式：
- .docx (Word 段落 + 表格)
- .pdf (pdfplumber)
- .pptx (python-pptx)
- .xlsx (openpyxl)
- .txt 编码自动检测
- 异常路径（不存在的文件 / 二进制兜底）
"""

import os
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

from scripts.file_importer import MultiFormatImporter


def _make_docx(path: str) -> None:
    try:
        from docx import Document
    except ImportError:
        raise unittest.SkipTest("python-docx not installed")
    doc = Document()
    doc.add_paragraph("小王: 你好")
    doc.add_paragraph("self: 我很好")
    doc.add_paragraph("今天吃饭")
    doc.add_paragraph("好的")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "表格发送者"
    tbl.cell(0, 1).text = "表格内容"
    tbl.cell(1, 0).text = "self"
    tbl.cell(1, 1).text = "嗯嗯"
    doc.save(path)


def _make_xlsx(path: str) -> None:
    try:
        import openpyxl
    except ImportError:
        raise unittest.SkipTest("openpyxl not installed")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "小王"
    ws["B1"] = "你好"
    ws["A2"] = "self"
    ws["B2"] = "我很好"
    ws["A3"] = "今天吃饭"
    ws["B3"] = "好的"
    wb.save(path)


def _make_pptx(path: str) -> None:
    try:
        from pptx import Presentation
    except ImportError:
        raise unittest.SkipTest("python-pptx not installed")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tx = slide.shapes.add_textbox(0, 0, 100, 100)
    tx.text_frame.text = "小王: PPT 内容"
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    tx2 = slide2.shapes.add_textbox(0, 0, 100, 100)
    tx2.text_frame.text = "self: 第二页内容"
    prs.save(path)


def _make_pdf(path: str) -> None:
    try:
        from reportlab.pdfgen import canvas
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        c = canvas.Canvas(path)
        c.setFont("STSong-Light", 14)
        c.drawString(100, 750, "测试 PDF 文档")
        c.drawString(100, 700, "小王: 你好")
        c.drawString(100, 680, "self: 我很好")
        c.drawString(100, 660, "今天吃饭")
        c.drawString(100, 640, "好的")
        c.showPage()
        c.save()
    except ImportError:
        # 跳过 PDF 测试
        raise unittest.SkipTest("reportlab not installed")


class TestFileImporterDocx(unittest.TestCase):
    """Word .docx 解析测试"""

    def setUp(self):
        self.imp = MultiFormatImporter({})
        self.tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
        self.tmp.close()

    def tearDown(self):
        try:
            os.unlink(self.tmp.name)
        except OSError:
            pass

    def test_docx_paragraphs(self):
        _make_docx(self.tmp.name)
        msgs = self.imp.import_file(self.tmp.name)
        # 至少 4 条普通段落消息
        self.assertGreaterEqual(len(msgs), 4)
        # 第一条应该是"小王: 你好"
        senders = [m.get("sender") for m in msgs]
        self.assertIn("小王", senders)
        self.assertIn("self", senders)

    def test_docx_table_content(self):
        _make_docx(self.tmp.name)
        msgs = self.imp.import_file(self.tmp.name)
        # 表格内容应包含
        contents = " ".join(m.get("content", "") for m in msgs)
        self.assertIn("表格", contents)

    def test_docx_metadata(self):
        _make_docx(self.tmp.name)
        msgs = self.imp.import_file(self.tmp.name)
        # 至少有一条带 doc_section metadata
        sections = [m.get("metadata", {}).get("doc_section") for m in msgs]
        self.assertIn("paragraph", sections)


class TestFileImporterXlsx(unittest.TestCase):
    """Excel .xlsx 解析测试"""

    def setUp(self):
        self.imp = MultiFormatImporter({})
        self.tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        self.tmp.close()

    def tearDown(self):
        try:
            os.unlink(self.tmp.name)
        except OSError:
            pass

    def test_xlsx_basic(self):
        _make_xlsx(self.tmp.name)
        msgs = self.imp.import_file(self.tmp.name)
        self.assertGreater(len(msgs), 0)
        # 至少包含一个 sender 标识
        all_text = " ".join(m.get("content", "") for m in msgs)
        self.assertIn("小王", all_text)
        self.assertIn("你好", all_text)

    def test_xlsx_metadata_sheet(self):
        _make_xlsx(self.tmp.name)
        msgs = self.imp.import_file(self.tmp.name)
        sheets = [m.get("metadata", {}).get("xlsx_sheet") for m in msgs]
        self.assertTrue(any(s for s in sheets if s))


class TestFileImporterPptx(unittest.TestCase):
    """PowerPoint .pptx 解析测试"""

    def setUp(self):
        self.imp = MultiFormatImporter({})
        self.tmp = tempfile.NamedTemporaryFile(suffix=".pptx", delete=False)
        self.tmp.close()

    def tearDown(self):
        try:
            os.unlink(self.tmp.name)
        except OSError:
            pass

    def test_pptx_basic(self):
        _make_pptx(self.tmp.name)
        msgs = self.imp.import_file(self.tmp.name)
        self.assertGreater(len(msgs), 0)
        # 至少包含两页
        slides = [m.get("metadata", {}).get("ppt_slide") for m in msgs]
        self.assertIn(1, slides)
        self.assertIn(2, slides)

    def test_pptx_content(self):
        _make_pptx(self.tmp.name)
        msgs = self.imp.import_file(self.tmp.name)
        all_text = " ".join(m.get("content", "") for m in msgs)
        self.assertIn("PPT", all_text)


class TestFileImporterPdf(unittest.TestCase):
    """PDF 解析测试"""

    def setUp(self):
        self.imp = MultiFormatImporter({})

    def test_pdf_basic(self):
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp:
            try:
                _make_pdf(tmp.name)
                msgs = self.imp.import_file(tmp.name)
                self.assertGreater(len(msgs), 0)
                # 至少识别出一个 sender
                senders = [m.get("sender") for m in msgs]
                self.assertTrue(any(s for s in senders if s and s != "document"))
            except unittest.SkipTest:
                self.skipTest("reportlab not available")
            finally:
                try:
                    os.unlink(tmp.name)
                except OSError:
                    pass


class TestFileImporterTxt(unittest.TestCase):
    """txt 编码自动检测测试"""

    def setUp(self):
        self.imp = MultiFormatImporter({})

    def test_txt_utf8(self):
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".txt", delete=False, encoding="utf-8"
        ) as tmp:
            tmp.write("小王: 你好\nself: 我很好\n")
            path = tmp.name
        try:
            msgs = self.imp.import_file(path)
            self.assertEqual(len(msgs), 2)
            self.assertEqual(msgs[0]["sender"], "小王")
        finally:
            os.unlink(path)

    def test_txt_gbk(self):
        data = "小王: 你好\nself: 我很好\n".encode("gb18030")
        with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tmp:
            tmp.write(data)
            path = tmp.name
        try:
            msgs = self.imp.import_file(path)
            self.assertEqual(len(msgs), 2)
            self.assertEqual(msgs[0]["sender"], "小王")
        finally:
            os.unlink(path)


class TestFileImporterErrors(unittest.TestCase):
    """异常路径测试"""

    def setUp(self):
        self.imp = MultiFormatImporter({})

    def test_missing_file(self):
        msgs = self.imp.import_file("/tmp/non_existent_file_xyz_docx_test.docx")
        self.assertEqual(msgs, [])

    def test_directory_path(self):
        """目录路径应返回 []"""
        msgs = self.imp.import_file("/tmp")
        self.assertEqual(msgs, [])

    def test_image_returns_image_msg(self):
        """图片应返回 image 类型消息"""
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp.write(b"fake png data")
            path = tmp.name
        try:
            msgs = self.imp.import_file(path)
            self.assertEqual(len(msgs), 1)
            self.assertEqual(msgs[0]["msg_type"], "image")
        finally:
            os.unlink(path)

    def test_audio_returns_audio_msg(self):
        with tempfile.NamedTemporaryFile(suffix=".mp3", delete=False) as tmp:
            tmp.write(b"fake mp3 data")
            path = tmp.name
        try:
            msgs = self.imp.import_file(path)
            self.assertEqual(len(msgs), 1)
            self.assertEqual(msgs[0]["msg_type"], "audio")
        finally:
            os.unlink(path)

    def test_video_returns_video_msg(self):
        with tempfile.NamedTemporaryFile(suffix=".mp4", delete=False) as tmp:
            tmp.write(b"fake mp4 data")
            path = tmp.name
        try:
            msgs = self.imp.import_file(path)
            self.assertEqual(len(msgs), 1)
            self.assertEqual(msgs[0]["msg_type"], "video")
        finally:
            os.unlink(path)


class TestFileImporterJson(unittest.TestCase):
    """JSON 解析测试"""

    def setUp(self):
        self.imp = MultiFormatImporter({})

    def test_json_list(self):
        import json
        data = [
            {"timestamp": "2024-01-01 12:00:00", "sender": "小王", "content": "你好"},
            {"timestamp": "2024-01-01 12:01:00", "sender": "self", "content": "我很好"},
        ]
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".json", delete=False, encoding="utf-8"
        ) as tmp:
            json.dump(data, tmp, ensure_ascii=False)
            path = tmp.name
        try:
            msgs = self.imp.import_file(path)
            self.assertEqual(len(msgs), 2)
            self.assertEqual(msgs[0]["sender"], "小王")
        finally:
            os.unlink(path)

    def test_json_messages_field(self):
        import json
        data = {
            "messages": [
                {"timestamp": "2024-01-01 12:00:00", "name": "小王", "message": "你好"},
            ]
        }
        with tempfile.NamedTemporaryFile(
            mode="w", suffix=".json", delete=False, encoding="utf-8"
        ) as tmp:
            json.dump(data, tmp, ensure_ascii=False)
            path = tmp.name
        try:
            msgs = self.imp.import_file(path)
            self.assertEqual(len(msgs), 1)
            self.assertEqual(msgs[0]["sender"], "小王")
        finally:
            os.unlink(path)


class TestExtractPrintableChinese(unittest.TestCase):
    """v2.5.0：二进制兜底的中文提取（原实现只保留 ASCII，中文全丢）"""

    def test_utf8_chinese_extracted(self):
        """UTF-8 编码的中文片段应被提取"""
        data = b"\x00\x01binary junk" + "小王: 你好\nself: 我很好".encode("utf-8") + b"\xff\xfe\x00"
        text = MultiFormatImporter._extract_printable_chinese(data)
        self.assertIn("你好", text)
        self.assertIn("小王", text)

    def test_gb18030_chinese_extracted(self):
        """GBK/GB18030 编码的中文也应被提取（多编码尝试）"""
        data = b"junk\x00" + "张三: 忙吗".encode("gb18030") + b"\x01\x02"
        text = MultiFormatImporter._extract_printable_chinese(data)
        self.assertIn("张三", text)

    def test_utf16le_chinese_extracted(self):
        """UTF-16LE（旧版 .doc 常见）中文提取"""
        data = b"\x00j\x00u\x00n\x00k" + "你好世界".encode("utf-16-le") + b"\x00\x00"
        text = MultiFormatImporter._extract_printable_chinese(data)
        self.assertIn("你好世界", text)

    def test_empty_input(self):
        self.assertEqual(MultiFormatImporter._extract_printable_chinese(b""), "")

    def test_pure_binary_returns_empty_or_clean(self):
        """纯二进制垃圾不产生可读行"""
        data = bytes(range(256)) * 2
        text = MultiFormatImporter._extract_printable_chinese(data)
        # 允许为空（乱码行被可读性过滤）
        self.assertNotIn("\ufffd\ufffd\ufffd\ufffd\ufffd", text)


if __name__ == "__main__":
    unittest.main()
