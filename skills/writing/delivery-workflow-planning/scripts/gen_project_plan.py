# -*- coding: utf-8 -*-
"""
生成《项目实施计划》（正排 / 倒排甘特表）

数据源: references/实施计划任务表.json（27 项任务标准参数，来自《实施方法论》ls3pcw）

排程逻辑：
- 正排:  --launch-date 项目启动日 → 按方法论任务表顺序向后推
        任务开始 = 前一任务完成日 + 前一任务 interval（阶段安排间隔周期）
- 倒排:  --online-date 上线日期   → 先按正排生成全表，再整体平移使 4.2 系统正式切换
        锚定上线日（推荐，保证间隔与《实施方法论》建议完全一致）
- 时数调整: --adjust "1.4:2,3.5:3" 按调研实施范围/功能范围调整指定任务辅导天数

用法:
  python gen_project_plan.py --client "客户名" --online-date 2026/08/01 --out plan.html
  python gen_project_plan.py --client "客户名" --online-date 2026/08/01 \
      --adjust "1.4:2,3.5:3" --out-format xlsx --out plan.xlsx
  python gen_project_plan.py --client "客户名" --launch-date 2026/06/01 --out plan.html
"""
import argparse
import json
import os
import sys
from datetime import date, datetime, timedelta

DATE_FMT = "%Y/%m/%d"

def parse_date(s):
    s = s.strip().replace("-", "/")
    return datetime.strptime(s, DATE_FMT).date()

def fmt(d):
    return d.strftime(DATE_FMT)

def load_tasks():
    here = os.path.dirname(os.path.abspath(__file__))
    ref = os.path.join(here, "..", "references", "实施计划任务表.json")
    with open(ref, encoding="utf-8") as f:
        data = json.load(f)
    return data["stages"]

def apply_adjust(stages, adjust_str):
    """--adjust "1.4:2,3.5:3"：覆盖指定任务辅导天数（按调研实施范围/功能范围）"""
    if not adjust_str:
        return
    for item in adjust_str.split(","):
        if ":" not in item:
            continue
        code, days = item.strip().split(":", 1)
        days = float(days.strip())
        found = False
        for st in stages:
            for t in st["tasks"]:
                if t["code"] == code.strip():
                    t["days"] = days
                    found = True
                    print(f"  调整 {code} 辅导天数 → {days} 天")
        if not found:
            print(f"  ⚠️ 未找到任务 {code}，忽略")

def schedule(stages, mode, anchor):
    """先按正排生成全部日期，再按模式处理锚点"""
    # 1. 正排生成
    rows_all = []  # (stage, task, start, end)
    cur = anchor if mode == "launch" else date(2000, 1, 1)  # 倒排先占位，稍后平移
    for st in stages:
        for t in st["tasks"]:
            start = cur
            end = start + timedelta(days=int(t["days"] or 0))
            rows_all.append((st["stage"], t, start, end))
            cur = end + timedelta(days=int(t["interval"] or 0))
    # 2. 倒排：平移使 4.2 锚定上线日
    if mode == "online":
        anchor42 = next((s for _, t, s, _ in rows_all if t["code"] == "4.2"), None)
        if anchor42 is None:
            raise RuntimeError("任务表中未找到 4.2 系统正式切换")
        delta = anchor - anchor42
        rows_all = [(stn, t, s + delta, e + delta) for stn, t, s, e in rows_all]
    # 3. 按阶段分组
    out = []
    cur_stage = None
    for (stn, t, s, e) in rows_all:
        if cur_stage != stn:
            out.append({"stage": stn, "rows": []})
            cur_stage = stn
        out[-1]["rows"].append({"t": t, "start": s, "end": e, "milestone": t["milestone"]})
    return out

def to_html(client, mode, anchor_str, stages_plan):
    rows_html = []
    milestones = []
    for st in stages_plan:
        rows = ""
        for r in st["rows"]:
            t = r["t"]
            if r["milestone"]:
                milestones.append((t["code"], t["name"], fmt(r["end"])))
            star = '<span style="color:#E6A23C">☆</span>' if r["milestone"] else ""
            rows += (
                f"<tr><td><b>{t['code']}</b> {star}</td><td>{t['name']}</td>"
                f"<td>{t['desc']}</td><td>{t['days']}</td>"
                f"<td>{fmt(r['start'])}</td><td>{fmt(r['end'])}</td></tr>"
            )
        rows_html.append(
            f"<h3>{st['stage']}</h3>"
            f"<table border='1' cellspacing='0' cellpadding='6' style='border-collapse:collapse;width:100%;font-size:13px'>"
            f"<tr style='background:#0B3D91;color:#fff'>"
            f"<th>编码</th><th>任务</th><th>说明</th><th>辅导(天)</th><th>开始</th><th>完成</th></tr>"
            f"{rows}</table>"
        )
    mile_html = "".join(
        f"<tr><td><b>{c}</b> ☆</td><td>{n}</td><td>{d}</td></tr>" for c, n, d in milestones
    )
    head = (
        f"<h1>{client} eMES 项目实施计划</h1>"
        f"<p>排程模式：{'正排（从项目启动日开始）' if mode=='launch' else '倒排（锚定上线日，按方法论间隔，推荐）'}　锚点：{anchor_str}</p>"
    )
    return (
        f"<html><head><meta charset='utf-8'><title>{client} 实施计划</title></head><body>"
        f"{head}<h2>里程碑一览（7 个 ☆）</h2>"
        f"<table border='1' cellspacing='0' cellpadding='6' style='border-collapse:collapse;width:60%;font-size:13px'>"
        f"<tr style='background:#0B3D91;color:#fff'><th>编码</th><th>里程碑</th><th>目标日期</th></tr>"
        f"{mile_html}</table>"
        f"<h2>分阶段计划</h2>{''.join(rows_html)}"
        f"<p style='color:#888;font-size:12px'>生成：智能交付助手-流程规划 · gen_project_plan.py　数据源：《实施方法论》任务表</p>"
        f"</body></html>"
    )

def to_xlsx(client, mode, anchor_str, stages_plan, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "实施计划"
    ws.append(["客户", client])
    ws.append(["排程模式", "倒排（锚定上线日）" if mode == "online" else "正排（从启动日开始）"])
    ws.append(["锚点", anchor_str])
    ws.append(["生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    ws.append([])

    hfill = PatternFill("solid", fgColor="0B3D91")
    mile_fill = PatternFill("solid", fgColor="FFF2CC")

    # 里程碑一览
    ws.append(["里程碑一览（7 个 ☆）"])
    ws.append(["编码", "里程碑", "目标日期"])
    for c in ws[ws.max_row]:
        c.font = Font(color="FFFFFF", bold=True)
        c.fill = hfill
    for st in stages_plan:
        for r in st["rows"]:
            if r["milestone"]:
                ws.append([r["t"]["code"], r["t"]["name"], fmt(r["end"])])
                for c in ws[ws.max_row]:
                    c.fill = mile_fill
    ws.append([])

    # 分阶段计划
    for st in stages_plan:
        ws.append([st["stage"]])
        ws.append(["编码", "任务", "说明", "辅导(天)", "开始", "完成", "里程碑"])
        for c in ws[ws.max_row]:
            c.font = Font(color="FFFFFF", bold=True)
            c.fill = hfill
        for r in st["rows"]:
            t = r["t"]
            ws.append([t["code"], t["name"], t["desc"], t["days"],
                       fmt(r["start"]), fmt(r["end"]), "☆" if r["milestone"] else ""])
        ws.append([])

    widths = {"A": 8, "B": 26, "C": 46, "D": 10, "E": 12, "F": 12, "G": 8}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--client", required=True)
    ap.add_argument("--launch-date", help="项目启动日 YYYY/MM/DD（正排）")
    ap.add_argument("--online-date", help="上线日期 YYYY/MM/DD（倒排，推荐）")
    ap.add_argument("--adjust", default="", help='按调研范围调整辅导天数，如 "1.4:2,3.5:3"')
    ap.add_argument("--out", default="项目实施计划.html")
    ap.add_argument("--out-format", default="html", choices=["html", "xlsx", "both"],
                    help="输出格式：html / xlsx / both")
    args = ap.parse_args()

    stages = load_tasks()
    apply_adjust(stages, args.adjust)
    if args.online_date:
        plan = schedule(stages, "online", parse_date(args.online_date))
        anchor_str = f"上线日 {args.online_date}"
        mode = "online"
    elif args.launch_date:
        plan = schedule(stages, "launch", parse_date(args.launch_date))
        anchor_str = f"启动日 {args.launch_date}"
        mode = "launch"
    else:
        print("错误：必须提供 --online-date 或 --launch-date 之一")
        return 1

    fmt_ = args.out_format
    base = os.path.splitext(args.out)[0]
    if fmt_ in ("html", "both"):
        with open(args.out if fmt_ == "html" else base + ".html", "w", encoding="utf-8") as f:
            f.write(to_html(args.client, mode, anchor_str, plan))
        print(f"已生成 HTML: {args.out if fmt_=='html' else base+'.html'}")
    if fmt_ in ("xlsx", "both"):
        out_x = args.out if fmt_ == "xlsx" else base + ".xlsx"
        to_xlsx(args.client, mode, anchor_str, plan, out_x)
        print(f"已生成 Excel: {out_x}")
    return 0

if __name__ == "__main__":
    sys.exit(main())
