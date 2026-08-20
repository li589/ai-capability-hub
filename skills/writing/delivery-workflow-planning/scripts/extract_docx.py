# -*- coding: utf-8 -*-
"""
DOCX 文本提取（依赖 python-docx）
- 提取段落 + 表格全部文本
- 用法: python extract_docx.py <file.docx> [输出.txt]
"""
import sys
import os


def extract_docx_text(path):
    """返回文档全文（段落 + 表格）"""
    from docx import Document
    doc = Document(path)
    out_lines = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            out_lines.append(t)
    for ti, table in enumerate(doc.tables):
        out_lines.append(f"\n[表格{ti+1}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            out_lines.append(" | ".join(cells))
    return "\n".join(out_lines)


def main():
    if len(sys.argv) < 2:
        print("用法: python extract_docx.py <file.docx> [输出.txt]")
        return 1
    path = sys.argv[1]
    text = extract_docx_text(path)
    out = sys.argv[2] if len(sys.argv) > 2 else path.replace(".docx", "_文本.txt")
    with open(out, "w", encoding="utf-8") as f:
        f.write(text)
    print(f"总字数: {len(text)}")
    print(text[:2000])
    return 0


if __name__ == "__main__":
    sys.exit(main())
