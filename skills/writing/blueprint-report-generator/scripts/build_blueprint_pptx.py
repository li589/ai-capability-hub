# -*- coding: utf-8 -*-
"""
操作流程界面补充页生成器（python-pptx 离线构造）
- 用途：slidep 实时编译撞 editor_sdk 文件锁时，用 python-pptx 绕开生成补充页
- 输入：resources/images/ 下的 ui_<子流程>_<终端>.png 截图（4 张推荐）
- 输出：<项目>_<客户>_操作流程界面补充页.pptx（5 页：封面 + N 张界面+说明）
- 依赖：python-pptx + Pillow
"""
import os
import sys
import argparse
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from PIL import Image

# 默认配色（南京鼎华模板 · 客户A V1 实测，与 SKILL.md §4.5 / references/design_template.md 一致）
NAVY = RGBColor(0x00, 0x7B, 0xD3)      # 深蓝（原 #0B5394）
CYAN = RGBColor(0x2D, 0xB8, 0xF1)      # 主蓝·鼎华标志色（原 #06B6D4）
LIGHT = RGBColor(0xF8, 0xF9, 0xFA)     # 页面底色（浅灰白）
DARK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xEE, 0x82, 0x2F)    # 琥珀橙（个案/待确认标注）


def add_rect(slide, x, y, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
    s.shadow.inherit = False
    return s


def add_text(slide, x, y, w, h, text, font_size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT, name='思源黑体'):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    lines = text.split('\n') if isinstance(text, str) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = name
    return tb


def fit_image(path, max_w_in, max_h_in):
    img = Image.open(path)
    iw, ih = img.size
    max_w_px = int(max_w_in * 96)
    max_h_px = int(max_h_in * 96)
    ratio = min(max_w_px / iw, max_h_px / ih, 1)
    return Inches(max_w_in * ratio), Inches(max_h_in * ratio)


def build_supplement_pptx(out_path, customer_name, cards):
    """cards: [{'img': 'ui_派工_平板.png', 'title': '① 任务聪明派 · 派工',
                  'who': '班组长 / 车间主任', 'plat': '平板 PADA_DISPATCH',
                  'screen': '待派工 / 派工操作 / 已派工汇总', 'flow': '对应 4.1 派工子流程'}, ...]
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    SLIDE_W = prs.slide_width
    SLIDE_H = prs.slide_height
    BLANK = prs.slide_layouts[6]

    # ============ 封面 ============
    slide = prs.slides.add_slide(BLANK)
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    # 装饰
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8), Inches(0.5), Inches(4.5), Inches(4.5))
    oval.fill.background(); oval.line.color.rgb = WHITE; oval.line.width = Pt(1); oval.shadow.inherit = False
    oval = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9), Inches(-0.5), Inches(6), Inches(6))
    oval.fill.background(); oval.line.color.rgb = WHITE; oval.line.width = Pt(0.5); oval.shadow.inherit = False

    add_rect(slide, Inches(0.9), Inches(2.55), Inches(0.6), Inches(0.07), CYAN)
    add_text(slide, Inches(0.9), Inches(2.8), Inches(8), Inches(0.5), customer_name, 22, False, RGBColor(0xCC, 0xDD, 0xEE))
    add_text(slide, Inches(0.9), Inches(3.6), Inches(11), Inches(1.3), 'eMES 蓝图规划报告 · 操作流程界面补充页', 38, True, WHITE)
    add_text(slide, Inches(0.9), Inches(4.8), Inches(11), Inches(0.6), '核心子流程 · 现场真实界面预览', 18, False, RGBColor(0xCC, 0xDD, 0xEE))
    add_text(slide, Inches(0.9), Inches(6.5), Inches(8), Inches(0.4), '鼎华智能', 13, False, RGBColor(0xAA, 0xCC, 0xEE))

    # ============ 每个卡片一页 ============
    n = len(cards)
    for i, c in enumerate(cards, 1):
        slide = prs.slides.add_slide(BLANK)
        add_rect(slide, 0, 0, SLIDE_W, Inches(0.9), NAVY)
        add_rect(slide, Inches(0.6), Inches(0.25), Inches(0.07), Inches(0.4), CYAN)
        add_text(slide, Inches(0.8), Inches(0.18), Inches(10), Inches(0.6), c['title'], 22, True, WHITE)
        add_text(slide, Inches(11), Inches(0.3), Inches(2), Inches(0.4), f'补充页 · 0{i}/{n:02d}', 11, False, RGBColor(0xCC, 0xDD, 0xEE), PP_ALIGN.RIGHT)

        add_rect(slide, 0, Inches(0.9), SLIDE_W, SLIDE_H - Inches(0.9), WHITE)

        # 左图卡
        card_x, card_y = Inches(0.5), Inches(1.1)
        card_w, card_h = Inches(7), Inches(5.8)
        add_rect(slide, card_x, card_y, card_w, card_h, LIGHT)
        img_path = c['img_path']
        if os.path.exists(img_path):
            iw, ih = fit_image(img_path, 6.7, 5.5)
            img_x = card_x + (card_w - iw) / 2
            img_y = card_y + (card_h - ih) / 2
            slide.shapes.add_picture(img_path, img_x, img_y, iw, ih)
        else:
            add_text(slide, card_x + Inches(0.5), card_y + Inches(2.5), Inches(6), Inches(0.5), f'[图片缺失: {c["img"]}]', 14, False, GREY, PP_ALIGN.CENTER)

        # 右说明
        desc_x, desc_y = Inches(7.8), Inches(1.1)
        desc_w = Inches(5.1)
        add_text(slide, desc_x, desc_y, desc_w, Inches(0.4), '操作人员', 12, True, GREY)
        add_text(slide, desc_x, desc_y + Inches(0.32), desc_w, Inches(0.5), c['who'], 18, False, DARK)
        add_text(slide, desc_x, desc_y + Inches(1.1), desc_w, Inches(0.4), '平台', 12, True, GREY)
        add_text(slide, desc_x, desc_y + Inches(1.42), desc_w, Inches(0.5), c['plat'], 18, False, DARK)
        add_text(slide, desc_x, desc_y + Inches(2.2), desc_w, Inches(0.4), '关键界面', 12, True, GREY)
        add_text(slide, desc_x, desc_y + Inches(2.52), desc_w, Inches(2.5), c['screen'], 16, False, DARK)
        add_rect(slide, desc_x, Inches(6.5), desc_w, Inches(0.02), RGBColor(0xE5, 0xE7, 0xEB))
        add_text(slide, desc_x, Inches(6.6), desc_w, Inches(0.4), c['flow'], 13, True, CYAN)

    prs.save(out_path)
    return len(prs.slides)


def main():
    parser = argparse.ArgumentParser(description="操作流程界面补充页生成器（绕开 slidep 文件锁）")
    parser.add_argument("--out", required=True, help="输出 pptx 路径")
    parser.add_argument("--customer", required=True, help="客户名称（封面标题用）")
    parser.add_argument("--img-dir", required=True, help="resources/images/ 目录")
    parser.add_argument("--cards", nargs='+', required=True,
                        help='卡片定义，格式：img=文件名,title=标题,who=操作人员,plat=平台,screen=关键界面,flow=对应流程（多张用空格分隔）')
    args = parser.parse_args()

    cards = []
    for raw in args.cards:
        item = {}
        for kv in raw.split(','):
            k, v = kv.split('=', 1)
            item[k.strip()] = v.strip()
        item['img_path'] = os.path.join(args.img_dir, item['img'])
        cards.append(item)

    n = build_supplement_pptx(args.out, args.customer, cards)
    print(f"[OK] 已生成: {args.out} ({n} 页)")


if __name__ == "__main__":
    main()