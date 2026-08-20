#!/usr/bin/env python3
"""RAGGED LEFT EDGE — two stacked blocks starting at almost, but not quite, the same x.

This check was nearly not shipped. The first framing flagged 10 of 12 slides on a delivered deck
and MISSED the ragged page a human reviewer had actually complained about; every guard below
exists because a specific class of false positive was measured and had to be removed. So this
suite is mostly negative: a card's inner label, a value label following its bar, two unrelated
elements that happen to share an x, a deck-wide designed indent, and a sub-perceptual 2px slip
must ALL be silent. A check that cries wolf on a correct deck is worse than no check, because it
teaches the reader to skip the output where the real findings live.

Each negative below names the real-deck case it came from.
"""
import io, pathlib, sys, tempfile

HERE = pathlib.Path(__file__).resolve().parent
SCRIPTS = HERE.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))

ok, bad = [], []
for _m in ("pptx",):
    try:
        __import__(_m)
    except ImportError:
        print(f"skip: {_m} not installed")
        raise SystemExit(0)

import deckkit as dk                                        # noqa: E402
import lint_deck as L                                       # noqa: E402
from pptx.dml.color import RGBColor as C                    # noqa: E402

dk.FONT = "Helvetica Neue"
TMP = pathlib.Path(tempfile.mkdtemp(prefix="ragged-"))
W, H = 10.0, 5.625
PAPER, INK = "F5F1E6", "1F3B2F"


def _t(s, x, y, w, h, txt, size=15, bold=False):
    dk.text(s, x, y, w, h, [[(txt, size, C.from_string(INK), bold, False)]])


def _lint(prs, name):
    p = TMP / name
    prs.save(str(p))
    import contextlib
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(buf):
        try:
            L.lint(str(p), static_ok=True)
        except SystemExit:
            pass
    return [l for l in buf.getvalue().splitlines() if "RAGGED LEFT EDGE" in l]


def _page(prs):
    s = dk.add_slide(prs)
    dk.slide_background(s, PAPER)
    return s


# ---------------------------------------------------------------- positive: it has power
prs = dk.blank_deck(W, H)
s = _page(prs)
_t(s, 0.80, 0.50, 8.4, 0.7, "A heading", 30, True)
_t(s, 0.88, 1.28, 8.4, 0.6, "the paragraph directly beneath it")
hits = _lint(prs, "misaligned.pptx")
if len(hits) == 1 and "8px" in hits[0]:
    ok.append("an 8px slip between a heading and the block beneath it IS reported")
else:
    bad.append(f"the 8px slip was not caught: {hits!r}")

# and it names both blocks and both x positions, so the fix is obvious without opening the file
if hits and "0.80in" in hits[0] and "0.88in" in hits[0] and "A heading" in hits[0]:
    ok.append("the message names both blocks and both x positions")
else:
    bad.append(f"the message does not locate the defect: {hits!r}")

# ---------------------------------------------------------------- negative controls
prs = dk.blank_deck(W, H)
s = _page(prs)
_t(s, 0.80, 0.50, 8.4, 0.7, "A heading", 30, True)
_t(s, 0.80, 1.28, 8.4, 0.6, "the paragraph directly beneath it")
if not _lint(prs, "aligned.pptx"):
    ok.append("a properly aligned pair is silent")
else:
    bad.append("an aligned pair was reported as ragged")

prs = dk.blank_deck(W, H)
s = _page(prs)
_t(s, 0.80, 0.50, 8.4, 0.7, "A heading", 30, True)
_t(s, 0.82, 1.28, 8.4, 0.6, "the paragraph directly beneath it")
if not _lint(prs, "twopx.pptx"):
    ok.append("a 2px slip is BELOW the band — invisible at projection size, so not reported "
              "(the delivered deck's only two survivors were exactly this)")
else:
    bad.append("a sub-perceptual 2px slip was reported")

prs = dk.blank_deck(W, H)
s = _page(prs)
_t(s, 0.80, 0.50, 8.4, 0.7, "A heading", 30, True)
_t(s, 1.20, 1.28, 8.0, 0.6, "a deliberately indented block")
if not _lint(prs, "indent.pptx"):
    ok.append("a 0.40in indent is ABOVE the band — it reads as a decision, not a wobble")
else:
    bad.append("a deliberate indent was reported as ragged")

# a label inside a card is indented on purpose (raw box lefts flagged 10 of 12 real slides)
prs = dk.blank_deck(W, H)
s = _page(prs)
dk.box(s, 0.80, 1.10, 4.0, 1.6, fill="FFFFFF")
_t(s, 0.80, 0.50, 8.4, 0.5, "A heading", 30, True)
_t(s, 0.88, 1.24, 3.6, 0.5, "a label inside the card")
if not _lint(prs, "card.pptx"):
    ok.append("a label nested inside a card is exempt — its indent is the card's padding")
else:
    bad.append("a card's inner label was compared against the heading outside it")

# a value label follows ITS bar: two bars of different length MUST give different lefts
prs = dk.blank_deck(W, H)
s = _page(prs)
dk.box(s, 2.00, 1.20, 3.10, 0.5, fill="C08A2E")
_t(s, 5.18, 1.24, 1.2, 0.4, "17.61%")
dk.box(s, 2.00, 1.85, 3.02, 0.5, fill="8A8377")
_t(s, 5.10, 1.89, 1.2, 0.4, "17.08%")
if not _lint(prs, "bars.pptx"):
    ok.append("two value labels trailing bars of different length are exempt — snapping them "
              "to one x would break the chart (the real deck's 17.61%/17.08% pair)")
else:
    bad.append("value labels following their own bars were reported as misaligned")

# two elements a reader never compares: far apart vertically
prs = dk.blank_deck(W, H)
s = _page(prs)
_t(s, 0.80, 0.50, 8.4, 0.7, "A heading", 30, True)
_t(s, 0.88, 3.60, 1.0, 0.4, "-0.6", 14)
if not _lint(prs, "unrelated.pptx"):
    ok.append("a heading and an annotation 3in below it are not compared (three of the first "
              "framing's false positives were exactly this)")
else:
    bad.append("two vertically unrelated elements were compared")

# side by side, not stacked
prs = dk.blank_deck(W, H)
s = _page(prs)
_t(s, 0.80, 1.50, 1.2, 0.5, "left col")
_t(s, 0.88, 1.50, 1.2, 0.5, "beside it")
if not _lint(prs, "sidebyside.pptx"):
    ok.append("two blocks on the same line are not judged as a vertical stack")
else:
    bad.append("side-by-side blocks were treated as stacked")

# a GROUP is an authored unit: a child's absolute x follows where the group was dropped, so it
# is not a decision about that child. OVERLAP already reasons this way and exempts same-group
# pairs; alignment needs the stronger form. Measured: two labels inside one composed diagram,
# 8px apart, were reported before this.
prs = dk.blank_deck(W, H)
s = _page(prs)
_t(s, 0.80, 1.00, 3.0, 0.4, "label A inside a diagram")
_t(s, 0.88, 1.45, 3.0, 0.4, "label B inside a diagram")
s.shapes.add_group_shape([x for x in list(s.shapes)])
if not _lint(prs, "grouped.pptx"):
    ok.append("two labels inside one GROUP are exempt — a group is an authored unit")
else:
    bad.append("a composed diagram's internal labels were reported as ragged")

prs = dk.blank_deck(W, H)                                   # …the identical pair, UNgrouped
s = _page(prs)
_t(s, 0.80, 1.00, 3.0, 0.4, "label A inside a diagram")
_t(s, 0.88, 1.45, 3.0, 0.4, "label B inside a diagram")
if len(_lint(prs, "ungrouped.pptx")) == 1:
    ok.append("the SAME pair ungrouped is still reported (the exemption is about grouping, "
              "not about the geometry)")
else:
    bad.append("the ungrouped control did not fire — the exemption is too wide")

# ---------------------------------------------------------------- deck-wide recurrence exemption
prs = dk.blank_deck(W, H)
for i in range(3):                                          # the same offset on every page
    s = _page(prs)
    _t(s, 0.88, 0.42, 3.0, 0.3, f"KICKER {i}", 10, True)
    _t(s, 0.80, 0.80, 8.4, 0.7, f"Section title {i}", 26, True)
hits = _lint(prs, "systematic.pptx")
if not hits:
    ok.append("an 8px kicker indent repeated on 3 pages is a DESIGN and stays silent (this was "
              "the loudest false signal on the delivered deck: slides 3, 6 and 9, identically)")
else:
    bad.append(f"a deck-wide designed offset was reported {len(hits)} time(s): {hits!r}")

prs = dk.blank_deck(W, H)
for i in range(3):                                          # …the same offset on ONE page only
    s = _page(prs)
    if i == 1:
        _t(s, 0.88, 0.42, 3.0, 0.3, "KICKER", 10, True)
    else:
        _t(s, 0.80, 0.42, 3.0, 0.3, "KICKER", 10, True)
    _t(s, 0.80, 0.80, 8.4, 0.7, f"Section title {i}", 26, True)
hits = _lint(prs, "oneoff.pptx")
if len(hits) == 1 and "slide 2" in hits[0]:
    ok.append("the SAME offset on one page only is a slip, and names that page")
else:
    bad.append(f"a one-off offset was not reported, or named the wrong slide: {hits!r}")

# ---------------------------------------------------------------- the band constants are shared
if (L._EDGE_LO, L._EDGE_HI) == (0.04, 0.12):
    ok.append("the reportable band is 0.04-0.12in, matching REGISTRATION DRIFT's 'a hair' band")
else:
    bad.append(f"the band drifted from the documented 0.04-0.12in: "
               f"{(L._EDGE_LO, L._EDGE_HI)}")

print("\n".join("  ok   " + x for x in ok))
if bad:
    print("\n".join("  FAIL " + x for x in bad))
print(f"\n{len(ok)} passed, {len(bad)} failed")
raise SystemExit(1 if bad else 0)
