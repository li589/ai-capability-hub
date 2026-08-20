"""智能交付助手-服务器配置评估：Word 文档样式函数库

生成 Word（.docx）输出时统一调用本模块的样式函数，保证所有交付文档风格一致：
- 主题：深蓝标题 / 红色警示 / 深蓝表头白字 / 斑马纹表格 / 绿✅琥珀⚠️结论标记
- 页眉页脚 + 页码

用法示例：
    from docx_style import (add_heading, add_table, add_note_box, add_badge_para,
                            add_mixed_para, add_para, add_header_footer, GREEN_OK, AMBER)
    doc = Document()
    ... 设置 sections ...
    add_header_footer(doc, '文档标题')
    add_badge_para(doc, '✅ 满足要求，可放心部署', 'E8F5E9', GREEN_OK)
    add_heading(doc, '一、部署方式速查', level=1)
    add_table(doc, ['参数', '配置'], [...])
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# ---------- 主题色 ----------
RED = RGBColor(0xC0, 0x1F, 0x2B)          # 警示红
TITLE_BLUE = RGBColor(0x1F, 0x4E, 0x79)   # 主标题深蓝
ACCENT = RGBColor(0x2E, 0x74, 0xB5)       # 强调蓝
DARK = RGBColor(0x2C, 0x2C, 0x2A)         # 正文深灰
MUTED = RGBColor(0x6B, 0x6B, 0x6B)        # 次要灰
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN_OK = RGBColor(0x1E, 0x7B, 0x34)     # 满足/达标
AMBER = RGBColor(0xB5, 0x6A, 0x00)        # 基本满足/需注意

# ---------- 底纹 ----------
HDR_FILL = '1F4E79'      # 表头深蓝
BAND_FILL = 'F2F7FB'     # 斑马纹
NOTE_FILL = 'FFF4E5'     # 黄色提醒底
WARN_FILL = 'FDECEA'     # 红色警告底
OK_FILL = 'E8F5E9'       # 绿色结论底
AMBER_FILL = 'FFF8E1'    # 琥珀提示底


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def set_cell_margins(cell, top=60, start=90, bottom=60, end=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for tag, val in (('w:top', top), ('w:start', start), ('w:bottom', bottom), ('w:end', end)):
        node = tcMar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tcMar.append(node)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table, color='C8D6E3', sz=6):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def set_table_width(table, width_cm=17.0):
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(int(width_cm * 567)))
    tblW.set(qn('w:type'), 'dxa')
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def add_run(p, text, bold=False, color=None, size=10.5):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    if color is not None:
        r.font.color.rgb = color
    return r


def set_para_format(p, before=0, after=6, indent=0, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if indent:
        pf.left_indent = Cm(indent)
    if align is not None:
        pf.alignment = align


def add_para(doc, text='', bold=False, color=None, size=10.5, before=0, after=6, indent=0, align=None):
    p = doc.add_paragraph()
    set_para_format(p, before=before, after=after, indent=indent, align=align)
    if text:
        add_run(p, text, bold=bold, color=color, size=size)
    return p


def add_mixed_para(doc, segments, size=10.5, before=0, after=6, indent=0, align=None):
    p = doc.add_paragraph()
    set_para_format(p, before=before, after=after, indent=indent, align=align)
    for seg in segments:
        text, color, bold = seg
        add_run(p, text, bold=bold, color=color, size=size)
    return p


def add_heading(doc, text, level=1):
    """level1：深蓝+底细线；level2：红色左侧色条；level3：深灰加粗"""
    p = doc.add_paragraph()
    if level == 1:
        add_run(p, text, bold=True, color=TITLE_BLUE, size=16)
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '12')
        bottom.set(qn('w:space'), '2')
        bottom.set(qn('w:color'), '2E74B5')
        pBdr.append(bottom)
        pPr.append(pBdr)
    elif level == 2:
        add_run(p, text, bold=True, color=ACCENT, size=13)
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '18')
        left.set(qn('w:space'), '4')
        left.set(qn('w:color'), 'C01F2B')
        pBdr.append(left)
        pPr.append(pBdr)
        p.paragraph_format.left_indent = Cm(0.15)
    else:
        add_run(p, text, bold=True, color=DARK, size=12)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    return p


def write_cell(p, val, default_size=9.5):
    """单元格内容：str | (text,color,bold) | [(text,color,bold),...]"""
    if isinstance(val, str):
        add_run(p, val, size=default_size)
        return
    if isinstance(val, tuple) and len(val) == 3:
        text, color, bold = val
        add_run(p, text, bold=bold, color=color, size=default_size)
        return
    if isinstance(val, list):
        for seg in val:
            if isinstance(seg, tuple) and len(seg) == 3:
                text, color, bold = seg
                add_run(p, text, bold=bold, color=color, size=default_size)
            elif isinstance(seg, str):
                add_run(p, seg, size=default_size)


def add_table(doc, headers, rows, band=True, font_size=9.5):
    """专业表格：深蓝表头白字 + 斑马纹 + 统一边框"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    set_table_borders(table)
    set_table_width(table)
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, HDR_FILL)
        set_cell_margins(cell)
        for p in cell.paragraphs:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
        p = cell.paragraphs[0]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_run(p, h, bold=True, color=WHITE, size=font_size)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            set_cell_margins(cell)
            if band and i % 2 == 1:
                set_cell_shading(cell, BAND_FILL)
            for p in cell.paragraphs:
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            write_cell(p, val, default_size=font_size)
    return table


def add_badge_para(doc, text, fill, color, bold=True, size=13):
    """结论横幅（绿色✅ / 琥珀⚠️）"""
    p = doc.add_paragraph()
    set_para_format(p, before=8, after=10)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    pPr.append(shd)
    pBdr = OxmlElement('w:pBdr')
    for edge in ('top', 'bottom', 'left', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '8')
        el.set(qn('w:space'), '4')
        el.set(qn('w:color'), '1E7B34')
        pBdr.append(el)
    pPr.append(pBdr)
    add_run(p, text, bold=bold, color=color, size=size)
    return p


def add_note_box(doc, segments, kind='note'):
    """提醒框：note=黄底，warn=红底。segments=[(text,color,bold),...]"""
    p = doc.add_paragraph()
    set_para_format(p, before=4, after=10)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    for edge in ('top', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '4')
        el.set(qn('w:color'), 'E5B96B' if kind == 'note' else 'D98880')
        pBdr.append(el)
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')
    left.set(qn('w:space'), '4')
    left.set(qn('w:color'), 'E5B96B' if kind == 'note' else 'C01F2B')
    pBdr.append(left)
    pPr.append(pBdr)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), NOTE_FILL if kind == 'note' else WARN_FILL)
    pPr.append(shd)
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.right_indent = Cm(0.2)
    for seg in segments:
        text, color, bold = seg
        add_run(p, text, bold=bold, color=color, size=10)


def add_header_footer(doc, header_title):
    """页眉（标题右对齐+细线）+ 页脚居中页码"""
    section = doc.sections[0]
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ''
    add_run(hp, header_title, bold=True, color=MUTED, size=8.5)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'B8C6D4')
    pBdr.append(bottom)
    pPr.append(pBdr)
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = ''
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def init_document(margins_cm=2.0):
    """创建标准文档：微软雅黑 + 页边距；返回 doc"""
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(10.5)
    style.font.color.rgb = DARK
    for section in doc.sections:
        section.left_margin = Cm(margins_cm)
        section.right_margin = Cm(margins_cm)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.header_distance = Cm(1.0)
        section.footer_distance = Cm(1.0)
    return doc
