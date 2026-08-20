#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""本地独立版能源日报。

特点：
- 仅使用本地 SQLite，默认不联网；
- 图形界面基于 Python 标准库 tkinter；
- 支持新增、编辑、删除、筛选、汇总、备份；
- 支持 CSV 导入导出；安装 openpyxl 后支持 XLSX 导入导出；
- 可直接运行，也可用 PyInstaller 打包成 EXE；
- 提供 --self-test 供交付验收。
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import shutil
import sqlite3
import sys
import tempfile
from contextlib import closing
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable, Sequence

APP_NAME = "PY能源日报"
APP_VERSION = "1.7.0"
ENERGY_TYPES = ("电", "水", "天然气", "蒸汽", "压缩空气", "其他")
STATUSES = ("草稿", "已提交", "已审核")
CSV_HEADERS = [
    "日期", "工厂", "车间", "班次", "能源类型", "表计名称", "期初读数", "期末读数", "倍率",
    "消耗量", "计量单位", "产量", "产量单位", "单耗", "目标单耗", "单耗偏差率", "单价",
    "成本", "状态", "备注",
]
FIELD_KEYS = [
    "report_date", "factory", "workshop", "shift_name", "energy_type", "meter_name",
    "opening_reading", "closing_reading", "multiplier", "consumption", "energy_unit",
    "output_qty", "output_unit", "unit_consumption", "target_unit_consumption",
    "unit_variance_rate", "unit_price", "cost", "status", "remark",
]
HEADER_ALIASES = {
    "日期": "report_date", "报表日期": "report_date", "report_date": "report_date",
    "工厂": "factory", "factory": "factory", "车间": "workshop", "workshop": "workshop",
    "班次": "shift_name", "shift": "shift_name", "shift_name": "shift_name",
    "能源类型": "energy_type", "能源": "energy_type", "energy_type": "energy_type",
    "表计名称": "meter_name", "表计": "meter_name", "meter_name": "meter_name",
    "期初读数": "opening_reading", "opening_reading": "opening_reading",
    "期末读数": "closing_reading", "closing_reading": "closing_reading",
    "倍率": "multiplier", "multiplier": "multiplier",
    "消耗量": "consumption", "consumption": "consumption",
    "计量单位": "energy_unit", "能源单位": "energy_unit", "energy_unit": "energy_unit",
    "产量": "output_qty", "output_qty": "output_qty", "产量单位": "output_unit", "output_unit": "output_unit",
    "单耗": "unit_consumption", "unit_consumption": "unit_consumption",
    "目标单耗": "target_unit_consumption", "target_unit_consumption": "target_unit_consumption",
    "单耗偏差率": "unit_variance_rate", "unit_variance_rate": "unit_variance_rate",
    "单价": "unit_price", "unit_price": "unit_price", "成本": "cost", "cost": "cost",
    "状态": "status", "status": "status", "备注": "remark", "remark": "remark",
}


def default_data_dir() -> Path:
    if os.name == "nt":
        base = Path(os.environ.get("LOCALAPPDATA") or Path.home())
    elif sys.platform == "darwin":
        base = Path.home() / "Library" / "Application Support"
    else:
        base = Path(os.environ.get("XDG_DATA_HOME") or (Path.home() / ".local" / "share"))
    return base / "PY能源日报"


def decimal_value(value: object, default: str = "0") -> Decimal:
    text = str(value if value is not None else "").strip().replace(",", "")
    if not text:
        text = default
    try:
        return Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"无法识别数字：{value!r}") from exc


def q(value: Decimal, places: str = "0.0001") -> Decimal:
    return value.quantize(Decimal(places), rounding=ROUND_HALF_UP)


def num_text(value: Decimal | float | int | str | None, places: str = "0.0001") -> str:
    number = q(decimal_value(value), places)
    text = format(number, "f")
    return text.rstrip("0").rstrip(".") if "." in text else text


def normalize_date(value: object) -> str:
    text = str(value or "").strip()
    if not text:
        raise ValueError("日期不能为空")
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y%m%d"):
        try:
            return dt.datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    try:
        return dt.date.fromisoformat(text[:10]).isoformat()
    except ValueError as exc:
        raise ValueError(f"日期格式错误：{text}，请使用 YYYY-MM-DD") from exc


def friendly_message(exc: BaseException) -> str:
    if isinstance(exc, FileNotFoundError):
        return "找不到文件，请确认路径和文件名后重试。"
    if isinstance(exc, PermissionError) or getattr(exc, "winerror", None) in {5, 32, 33}:
        return "文件可能被 Excel、WPS 或同步软件占用，请关闭占用程序后重试。"
    if isinstance(exc, UnicodeError):
        return "文件编码无法识别，请另存为 UTF-8、UTF-8 BOM 或 GB18030。"
    if isinstance(exc, sqlite3.DatabaseError):
        return "数据库无法安全读取或写入，请停止操作并从备份恢复。"
    if isinstance(exc, ValueError):
        return str(exc).splitlines()[0]
    return "操作未完成，现有数据未被删除。请检查输入文件和目录权限后重试。"


def compute_metrics(payload: dict[str, object]) -> dict[str, str]:
    opening = decimal_value(payload.get("opening_reading"))
    closing = decimal_value(payload.get("closing_reading"))
    multiplier = decimal_value(payload.get("multiplier"), "1")
    if multiplier <= 0:
        raise ValueError("倍率必须大于 0")
    if closing < opening:
        raise ValueError("期末读数不能小于期初读数")
    consumption = q((closing - opening) * multiplier)
    output_qty = decimal_value(payload.get("output_qty"))
    unit_price = decimal_value(payload.get("unit_price"))
    target = decimal_value(payload.get("target_unit_consumption"))
    unit_consumption = q(consumption / output_qty) if output_qty > 0 else Decimal("0")
    variance = q((unit_consumption - target) / target * Decimal("100")) if target > 0 else Decimal("0")
    cost = q(consumption * unit_price, "0.01")
    return {
        "opening_reading": num_text(opening),
        "closing_reading": num_text(closing),
        "multiplier": num_text(multiplier),
        "consumption": num_text(consumption),
        "output_qty": num_text(output_qty),
        "unit_consumption": num_text(unit_consumption),
        "target_unit_consumption": num_text(target),
        "unit_variance_rate": num_text(variance),
        "unit_price": num_text(unit_price),
        "cost": num_text(cost, "0.01"),
    }


@dataclass(frozen=True)
class FilterSpec:
    date_from: str = ""
    date_to: str = ""
    factory: str = ""
    energy_type: str = ""
    status: str = ""
    keyword: str = ""


class EnergyRepository:
    def __init__(self, db_path: Path):
        self.db_path = db_path.expanduser().resolve()
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.initialize()

    def connect(self) -> sqlite3.Connection:
        con = sqlite3.connect(self.db_path)
        con.row_factory = sqlite3.Row
        con.execute("PRAGMA foreign_keys=ON")
        con.execute("PRAGMA busy_timeout=5000")
        return con

    def initialize(self) -> None:
        with closing(self.connect()) as con, con:
            con.executescript(
                """
                PRAGMA journal_mode=WAL;
                CREATE TABLE IF NOT EXISTS energy_daily_records (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    report_date TEXT NOT NULL,
                    factory TEXT NOT NULL,
                    workshop TEXT NOT NULL DEFAULT '',
                    shift_name TEXT NOT NULL DEFAULT '',
                    energy_type TEXT NOT NULL,
                    meter_name TEXT NOT NULL,
                    opening_reading TEXT NOT NULL DEFAULT '0',
                    closing_reading TEXT NOT NULL DEFAULT '0',
                    multiplier TEXT NOT NULL DEFAULT '1',
                    consumption TEXT NOT NULL DEFAULT '0',
                    energy_unit TEXT NOT NULL DEFAULT '',
                    output_qty TEXT NOT NULL DEFAULT '0',
                    output_unit TEXT NOT NULL DEFAULT '',
                    unit_consumption TEXT NOT NULL DEFAULT '0',
                    target_unit_consumption TEXT NOT NULL DEFAULT '0',
                    unit_variance_rate TEXT NOT NULL DEFAULT '0',
                    unit_price TEXT NOT NULL DEFAULT '0',
                    cost TEXT NOT NULL DEFAULT '0',
                    status TEXT NOT NULL DEFAULT '草稿',
                    remark TEXT NOT NULL DEFAULT '',
                    created_at TEXT NOT NULL,
                    updated_at TEXT NOT NULL,
                    UNIQUE(report_date, factory, workshop, shift_name, energy_type, meter_name)
                );
                CREATE INDEX IF NOT EXISTS idx_energy_daily_date ON energy_daily_records(report_date);
                CREATE INDEX IF NOT EXISTS idx_energy_daily_factory ON energy_daily_records(factory);
                CREATE INDEX IF NOT EXISTS idx_energy_daily_type ON energy_daily_records(energy_type);
                CREATE TABLE IF NOT EXISTS audit_log (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    action TEXT NOT NULL,
                    record_id INTEGER,
                    detail_json TEXT NOT NULL,
                    created_at TEXT NOT NULL
                );
                """
            )

    @staticmethod
    def normalize_payload(payload: dict[str, object]) -> dict[str, str]:
        data = {key: str(payload.get(key, "") or "").strip() for key in FIELD_KEYS}
        data["report_date"] = normalize_date(data["report_date"])
        for key, label in (("factory", "工厂"), ("energy_type", "能源类型"), ("meter_name", "表计名称")):
            if not data[key]:
                raise ValueError(f"{label}不能为空")
        if data["energy_type"] not in ENERGY_TYPES:
            data["energy_type"] = data["energy_type"] or "其他"
        data["status"] = data["status"] if data["status"] in STATUSES else "草稿"
        data.update(compute_metrics(data))
        return data

    def save(self, payload: dict[str, object], record_id: int | None = None) -> int:
        data = self.normalize_payload(payload)
        now = dt.datetime.now().isoformat(timespec="seconds")
        columns = FIELD_KEYS
        with closing(self.connect()) as con, con:
            try:
                if record_id is None:
                    sql = f"INSERT INTO energy_daily_records ({','.join(columns)},created_at,updated_at) VALUES ({','.join('?' for _ in columns)},?,?)"
                    cur = con.execute(sql, [data[key] for key in columns] + [now, now])
                    record_id = int(cur.lastrowid)
                    action = "create"
                else:
                    existing = con.execute("SELECT * FROM energy_daily_records WHERE id=?", (record_id,)).fetchone()
                    if existing is None:
                        raise ValueError("要编辑的记录不存在")
                    assignments = ",".join(f"{key}=?" for key in columns)
                    con.execute(f"UPDATE energy_daily_records SET {assignments},updated_at=? WHERE id=?", [data[key] for key in columns] + [now, record_id])
                    action = "update"
                con.execute(
                    "INSERT INTO audit_log(action,record_id,detail_json,created_at) VALUES (?,?,?,?)",
                    (action, record_id, json.dumps(data, ensure_ascii=False), now),
                )
                return int(record_id)
            except sqlite3.IntegrityError as exc:
                raise ValueError("同一日期、工厂、车间、班次、能源类型和表计名称的记录已存在") from exc

    def delete(self, record_id: int) -> None:
        now = dt.datetime.now().isoformat(timespec="seconds")
        with closing(self.connect()) as con, con:
            row = con.execute("SELECT * FROM energy_daily_records WHERE id=?", (record_id,)).fetchone()
            if row is None:
                return
            con.execute("DELETE FROM energy_daily_records WHERE id=?", (record_id,))
            con.execute(
                "INSERT INTO audit_log(action,record_id,detail_json,created_at) VALUES (?,?,?,?)",
                ("delete", record_id, json.dumps(dict(row), ensure_ascii=False), now),
            )

    def get(self, record_id: int) -> dict[str, object] | None:
        with closing(self.connect()) as con, con:
            row = con.execute("SELECT * FROM energy_daily_records WHERE id=?", (record_id,)).fetchone()
            return dict(row) if row else None

    def list(self, spec: FilterSpec = FilterSpec()) -> list[dict[str, object]]:
        where: list[str] = []
        params: list[str] = []
        if spec.date_from:
            where.append("report_date>=?")
            params.append(normalize_date(spec.date_from))
        if spec.date_to:
            where.append("report_date<=?")
            params.append(normalize_date(spec.date_to))
        if spec.factory:
            where.append("factory LIKE ?")
            params.append(f"%{spec.factory.strip()}%")
        if spec.energy_type:
            where.append("energy_type=?")
            params.append(spec.energy_type)
        if spec.status:
            where.append("status=?")
            params.append(spec.status)
        if spec.keyword:
            where.append("(workshop LIKE ? OR meter_name LIKE ? OR remark LIKE ?)")
            params.extend([f"%{spec.keyword.strip()}%"] * 3)
        clause = " WHERE " + " AND ".join(where) if where else ""
        sql = "SELECT * FROM energy_daily_records" + clause + " ORDER BY report_date DESC,factory,workshop,energy_type,meter_name"
        with closing(self.connect()) as con, con:
            return [dict(row) for row in con.execute(sql, params).fetchall()]

    def summary(self, spec: FilterSpec = FilterSpec()) -> list[dict[str, str]]:
        rows = self.list(spec)
        groups: dict[str, dict[str, Decimal]] = {}
        for row in rows:
            key = str(row["energy_type"])
            item = groups.setdefault(key, {"consumption": Decimal("0"), "cost": Decimal("0"), "output": Decimal("0")})
            item["consumption"] += decimal_value(row["consumption"])
            item["cost"] += decimal_value(row["cost"])
            item["output"] += decimal_value(row["output_qty"])
        result = []
        for energy_type, item in sorted(groups.items()):
            unit_consumption = item["consumption"] / item["output"] if item["output"] > 0 else Decimal("0")
            result.append({
                "energy_type": energy_type,
                "consumption": num_text(item["consumption"]),
                "cost": num_text(item["cost"], "0.01"),
                "output": num_text(item["output"]),
                "unit_consumption": num_text(unit_consumption),
            })
        return result

    def backup(self, target: Path) -> Path:
        target = target.expanduser().resolve()
        target.parent.mkdir(parents=True, exist_ok=True)
        with closing(self.connect()) as source, closing(sqlite3.connect(target)) as destination:
            source.backup(destination)
        return target

    def integrity_check(self) -> str:
        with closing(self.connect()) as con, con:
            return str(con.execute("PRAGMA integrity_check").fetchone()[0])


def rows_from_csv(path: Path) -> list[dict[str, object]]:
    errors: list[UnicodeDecodeError] = []
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                return [dict(row) for row in csv.DictReader(handle)]
        except UnicodeDecodeError as exc:
            errors.append(exc)
    raise ValueError(f"CSV 编码无法识别：{errors[-1] if errors else path}")


def rows_from_xlsx(path: Path) -> list[dict[str, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 XLSX 需要 openpyxl；请由 IT 使用已校验的离线依赖包安装") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    try:
        headers = [str(item or "").strip() for item in next(values)]
    except StopIteration:
        return []
    return [{headers[index]: value for index, value in enumerate(row) if index < len(headers)} for row in values]


def map_import_row(row: dict[str, object]) -> dict[str, object]:
    mapped: dict[str, object] = {}
    for header, value in row.items():
        key = HEADER_ALIASES.get(str(header or "").strip())
        if key:
            mapped[key] = value
    mapped.setdefault("multiplier", "1")
    mapped.setdefault("status", "草稿")
    mapped.setdefault("energy_unit", "")
    mapped.setdefault("output_unit", "")
    return mapped


def import_file(repository: EnergyRepository, path: Path) -> dict[str, object]:
    path = path.expanduser().resolve()
    if path.suffix.casefold() == ".csv":
        rows = rows_from_csv(path)
    elif path.suffix.casefold() in {".xlsx", ".xlsm"}:
        rows = rows_from_xlsx(path)
    else:
        raise ValueError("仅支持 CSV、XLSX、XLSM")
    success = 0
    failed: list[dict[str, object]] = []
    for index, row in enumerate(rows, start=2):
        try:
            repository.save(map_import_row(row))
            success += 1
        except Exception as exc:
            failed.append({"row": index, "error": friendly_message(exc)})
    return {"total": len(rows), "success": success, "failed": failed}


def export_csv(rows: Sequence[dict[str, object]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(CSV_HEADERS)
        for row in rows:
            writer.writerow([row.get(key, "") for key in FIELD_KEYS])


def export_xlsx(rows: Sequence[dict[str, object]], path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RuntimeError("导出 XLSX 需要 openpyxl；请由 IT 使用已校验的离线依赖包安装") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "能源日报"
    sheet.append(CSV_HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(key, "") for key in FIELD_KEYS])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [12, 14, 14, 10, 12, 18, 12, 12, 8, 12, 10, 12, 10, 12, 12, 12, 10, 12, 10, 24]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def export_file(repository: EnergyRepository, path: Path, spec: FilterSpec = FilterSpec()) -> int:
    rows = repository.list(spec)
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        export_csv(rows, path)
    elif suffix == ".xlsx":
        export_xlsx(rows, path)
    else:
        raise ValueError("导出文件扩展名必须是 .csv 或 .xlsx")
    return len(rows)


def run_self_test(db_path: Path | None = None) -> dict[str, object]:
    temporary = None
    if db_path is None:
        temporary = tempfile.TemporaryDirectory()
        db_path = Path(temporary.name) / "self_test.db"
    repo = EnergyRepository(db_path)
    payload = {
        "report_date": "2026-08-06", "factory": "测试工厂", "workshop": "动力车间", "shift_name": "白班",
        "energy_type": "电", "meter_name": "总表", "opening_reading": "100", "closing_reading": "125.5",
        "multiplier": "2", "energy_unit": "kWh", "output_qty": "1000", "output_unit": "吨",
        "target_unit_consumption": "0.04", "unit_price": "0.8", "status": "草稿", "remark": "自检数据",
    }
    record_id = repo.save(payload)
    row = repo.get(record_id)
    assert row is not None
    assert row["consumption"] == "51"
    assert row["cost"] == "40.8"
    repo.save({**payload, "closing_reading": "130"}, record_id=record_id)
    assert repo.get(record_id)["consumption"] == "60"
    summary = repo.summary()
    assert summary and summary[0]["energy_type"] == "电"
    csv_path = db_path.parent / "self_test.csv"
    assert export_file(repo, csv_path) == 1
    backup_path = db_path.parent / "self_test_backup.db"
    repo.backup(backup_path)
    assert backup_path.is_file()
    assert repo.integrity_check() == "ok"
    repo.delete(record_id)
    assert repo.list() == []
    result = {"ok": True, "version": APP_VERSION, "database": str(db_path), "checks": 9}
    if temporary is not None:
        temporary.cleanup()
    return result


class EnergyDailyApp:
    def __init__(self, root, repository: EnergyRepository):
        import tkinter as tk
        from tkinter import ttk

        self.tk = tk
        self.ttk = ttk
        self.root = root
        self.repo = repository
        self.root.title(f"{APP_NAME} V{APP_VERSION}")
        self.root.geometry("1460x820")
        self.root.minsize(1180, 680)
        self.status_var = tk.StringVar(value=f"数据库：{self.repo.db_path}")
        self.filter_vars = {
            "date_from": tk.StringVar(), "date_to": tk.StringVar(), "factory": tk.StringVar(),
            "energy_type": tk.StringVar(), "status": tk.StringVar(), "keyword": tk.StringVar(),
        }
        self._build_ui()
        self.refresh()

    def _build_ui(self) -> None:
        from tkinter import ttk

        top = ttk.Frame(self.root, padding=8)
        top.pack(fill="x")
        labels = [("起始日期", "date_from", 12), ("结束日期", "date_to", 12), ("工厂", "factory", 12), ("能源类型", "energy_type", 10), ("状态", "status", 10), ("关键词", "keyword", 16)]
        for label, key, width in labels:
            ttk.Label(top, text=label).pack(side="left", padx=(0, 3))
            if key == "energy_type":
                widget = ttk.Combobox(top, textvariable=self.filter_vars[key], values=("",) + ENERGY_TYPES, width=width, state="readonly")
            elif key == "status":
                widget = ttk.Combobox(top, textvariable=self.filter_vars[key], values=("",) + STATUSES, width=width, state="readonly")
            else:
                widget = ttk.Entry(top, textvariable=self.filter_vars[key], width=width)
            widget.pack(side="left", padx=(0, 8))
        ttk.Button(top, text="查询", command=self.refresh).pack(side="left", padx=3)
        ttk.Button(top, text="重置", command=self.reset_filters).pack(side="left", padx=3)

        toolbar = ttk.Frame(self.root, padding=(8, 0, 8, 8))
        toolbar.pack(fill="x")
        for text, command in [
            ("新增", self.add_record), ("编辑", self.edit_record), ("复制", self.duplicate_record),
            ("删除", self.delete_record), ("导入", self.import_dialog), ("导出", self.export_dialog),
            ("汇总", self.show_summary), ("备份", self.backup_dialog), ("刷新", self.refresh),
        ]:
            ttk.Button(toolbar, text=text, command=command).pack(side="left", padx=3)

        frame = ttk.Frame(self.root, padding=(8, 0, 8, 0))
        frame.pack(fill="both", expand=True)
        columns = ["id", "report_date", "factory", "workshop", "shift_name", "energy_type", "meter_name", "consumption", "energy_unit", "output_qty", "unit_consumption", "target_unit_consumption", "unit_variance_rate", "cost", "status", "remark"]
        headings = ["ID", "日期", "工厂", "车间", "班次", "能源类型", "表计", "消耗量", "单位", "产量", "单耗", "目标单耗", "偏差率%", "成本", "状态", "备注"]
        self.tree = ttk.Treeview(frame, columns=columns, show="headings", selectmode="browse")
        widths = [55, 95, 120, 120, 75, 85, 130, 90, 65, 85, 85, 85, 85, 90, 75, 220]
        for col, heading, width in zip(columns, headings, widths):
            self.tree.heading(col, text=heading, command=lambda c=col: self.sort_tree(c, False))
            self.tree.column(col, width=width, minwidth=50, anchor="center" if col != "remark" else "w")
        yscroll = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        xscroll = ttk.Scrollbar(frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)
        self.tree.bind("<Double-1>", lambda _event: self.edit_record())
        ttk.Label(self.root, textvariable=self.status_var, relief="sunken", anchor="w", padding=5).pack(fill="x", side="bottom")

    def filter_spec(self) -> FilterSpec:
        return FilterSpec(**{key: var.get().strip() for key, var in self.filter_vars.items()})

    def reset_filters(self) -> None:
        for var in self.filter_vars.values():
            var.set("")
        self.refresh()

    def refresh(self) -> None:
        from tkinter import messagebox
        try:
            rows = self.repo.list(self.filter_spec())
        except Exception as exc:
            messagebox.showerror("查询失败", friendly_message(exc), parent=self.root)
            return
        self.tree.delete(*self.tree.get_children())
        for row in rows:
            values = [row.get(col, "") for col in self.tree["columns"]]
            self.tree.insert("", "end", iid=str(row["id"]), values=values)
        total_cost = sum(decimal_value(row["cost"]) for row in rows)
        self.status_var.set(f"共 {len(rows)} 条 | 合计成本 {num_text(total_cost, '0.01')} | 数据库：{self.repo.db_path}")

    def selected_id(self) -> int | None:
        selection = self.tree.selection()
        return int(selection[0]) if selection else None

    def add_record(self) -> None:
        self.open_form(None)

    def edit_record(self) -> None:
        from tkinter import messagebox
        record_id = self.selected_id()
        if record_id is None:
            messagebox.showinfo("提示", "请先选择一条记录", parent=self.root)
            return
        self.open_form(record_id)

    def duplicate_record(self) -> None:
        from tkinter import messagebox
        record_id = self.selected_id()
        if record_id is None:
            messagebox.showinfo("提示", "请先选择一条记录", parent=self.root)
            return
        data = self.repo.get(record_id) or {}
        data["meter_name"] = str(data.get("meter_name", "")) + "-复制"
        self.open_form(None, data)

    def delete_record(self) -> None:
        from tkinter import messagebox
        record_id = self.selected_id()
        if record_id is None:
            messagebox.showinfo("提示", "请先选择一条记录", parent=self.root)
            return
        if messagebox.askyesno("确认删除", "删除后仍会保留审计记录，确定继续？", parent=self.root):
            self.repo.delete(record_id)
            self.refresh()

    def open_form(self, record_id: int | None, initial: dict[str, object] | None = None) -> None:
        import tkinter as tk
        from tkinter import messagebox, ttk
        data = initial or (self.repo.get(record_id) if record_id else {}) or {}
        win = tk.Toplevel(self.root)
        win.title("编辑能源日报" if record_id else "新增能源日报")
        win.transient(self.root)
        win.grab_set()
        fields = [
            ("日期*", "report_date", "entry", dt.date.today().isoformat()), ("工厂*", "factory", "entry", ""),
            ("车间", "workshop", "entry", ""), ("班次", "shift_name", "entry", ""),
            ("能源类型*", "energy_type", "energy", "电"), ("表计名称*", "meter_name", "entry", ""),
            ("期初读数", "opening_reading", "entry", "0"), ("期末读数", "closing_reading", "entry", "0"),
            ("倍率", "multiplier", "entry", "1"), ("计量单位", "energy_unit", "entry", "kWh"),
            ("产量", "output_qty", "entry", "0"), ("产量单位", "output_unit", "entry", "吨"),
            ("目标单耗", "target_unit_consumption", "entry", "0"), ("单价", "unit_price", "entry", "0"),
            ("状态", "status", "status", "草稿"), ("备注", "remark", "entry", ""),
        ]
        vars_: dict[str, tk.StringVar] = {}
        for index, (label, key, kind, default) in enumerate(fields):
            row, col = divmod(index, 2)
            base_col = col * 2
            ttk.Label(win, text=label).grid(row=row, column=base_col, padx=8, pady=6, sticky="e")
            var = tk.StringVar(value=str(data.get(key, default) or default))
            vars_[key] = var
            if kind == "energy":
                widget = ttk.Combobox(win, textvariable=var, values=ENERGY_TYPES, state="readonly", width=25)
            elif kind == "status":
                widget = ttk.Combobox(win, textvariable=var, values=STATUSES, state="readonly", width=25)
            else:
                widget = ttk.Entry(win, textvariable=var, width=28)
            widget.grid(row=row, column=base_col + 1, padx=8, pady=6, sticky="ew")
        preview = tk.StringVar(value="消耗量、单耗、偏差率和成本将在保存时自动计算")
        ttk.Label(win, textvariable=preview).grid(row=8, column=0, columnspan=4, padx=8, pady=(8, 2), sticky="w")

        def update_preview(*_args) -> None:
            try:
                metrics = compute_metrics({key: var.get() for key, var in vars_.items()})
                preview.set(f"消耗量 {metrics['consumption']} | 单耗 {metrics['unit_consumption']} | 偏差率 {metrics['unit_variance_rate']}% | 成本 {metrics['cost']}")
            except Exception as exc:
                preview.set(friendly_message(exc))

        for key in ("opening_reading", "closing_reading", "multiplier", "output_qty", "target_unit_consumption", "unit_price"):
            vars_[key].trace_add("write", update_preview)
        update_preview()

        def submit() -> None:
            try:
                self.repo.save({key: var.get() for key, var in vars_.items()}, record_id=record_id)
            except Exception as exc:
                messagebox.showerror("保存失败", friendly_message(exc), parent=win)
                return
            win.destroy()
            self.refresh()

        buttons = ttk.Frame(win)
        buttons.grid(row=9, column=0, columnspan=4, pady=12)
        ttk.Button(buttons, text="保存", command=submit).pack(side="left", padx=6)
        ttk.Button(buttons, text="取消", command=win.destroy).pack(side="left", padx=6)
        for col in range(4):
            win.columnconfigure(col, weight=1 if col % 2 else 0)

    def import_dialog(self) -> None:
        from tkinter import filedialog, messagebox
        path = filedialog.askopenfilename(parent=self.root, title="导入能源日报", filetypes=[("支持文件", "*.csv *.xlsx *.xlsm"), ("CSV", "*.csv"), ("Excel", "*.xlsx *.xlsm")])
        if not path:
            return
        try:
            result = import_file(self.repo, Path(path))
        except Exception as exc:
            messagebox.showerror("导入失败", friendly_message(exc), parent=self.root)
            return
        self.refresh()
        detail = "\n".join(f"第 {item['row']} 行：{item['error']}" for item in result["failed"][:10])
        messagebox.showinfo("导入完成", f"总行数：{result['total']}\n成功：{result['success']}\n失败：{len(result['failed'])}" + (f"\n\n{detail}" if detail else ""), parent=self.root)

    def export_dialog(self) -> None:
        from tkinter import filedialog, messagebox
        path = filedialog.asksaveasfilename(parent=self.root, title="导出能源日报", defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx"), ("CSV", "*.csv")])
        if not path:
            return
        try:
            count = export_file(self.repo, Path(path), self.filter_spec())
        except Exception as exc:
            messagebox.showerror("导出失败", friendly_message(exc), parent=self.root)
            return
        messagebox.showinfo("导出完成", f"已导出 {count} 条记录\n{path}", parent=self.root)

    def backup_dialog(self) -> None:
        from tkinter import filedialog, messagebox
        default_name = f"能源日报备份_{dt.datetime.now():%Y%m%d_%H%M%S}.db"
        path = filedialog.asksaveasfilename(parent=self.root, title="备份数据库", initialfile=default_name, defaultextension=".db", filetypes=[("SQLite 数据库", "*.db")])
        if not path:
            return
        try:
            self.repo.backup(Path(path))
        except Exception as exc:
            messagebox.showerror("备份失败", friendly_message(exc), parent=self.root)
            return
        messagebox.showinfo("备份完成", path, parent=self.root)

    def show_summary(self) -> None:
        import tkinter as tk
        from tkinter import ttk
        rows = self.repo.summary(self.filter_spec())
        win = tk.Toplevel(self.root)
        win.title("能源汇总")
        win.geometry("760x420")
        columns = ("energy_type", "consumption", "cost", "output", "unit_consumption")
        headings = ("能源类型", "总消耗量", "总成本", "总产量", "综合单耗")
        tree = ttk.Treeview(win, columns=columns, show="headings")
        for col, heading in zip(columns, headings):
            tree.heading(col, text=heading)
            tree.column(col, width=140, anchor="center")
        for row in rows:
            tree.insert("", "end", values=[row[col] for col in columns])
        tree.pack(fill="both", expand=True, padx=10, pady=10)
        ttk.Button(win, text="关闭", command=win.destroy).pack(pady=(0, 10))

    def sort_tree(self, column: str, reverse: bool) -> None:
        items = [(self.tree.set(item, column), item) for item in self.tree.get_children("")]
        def key(pair):
            value = pair[0]
            try:
                return (0, Decimal(value))
            except Exception:
                return (1, value)
        items.sort(key=key, reverse=reverse)
        for index, (_value, item) in enumerate(items):
            self.tree.move(item, "", index)
        self.tree.heading(column, command=lambda: self.sort_tree(column, not reverse))


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="本地独立版能源日报")
    parser.add_argument("--db", type=Path, default=default_data_dir() / "energy_daily.db", help="SQLite 数据库路径")
    parser.add_argument("--import-file", type=Path, help="导入 CSV/XLSX/XLSM")
    parser.add_argument("--export-file", type=Path, help="导出 CSV/XLSX")
    parser.add_argument("--no-gui", action="store_true", help="不启动图形界面")
    parser.add_argument("--self-test", action="store_true", help="运行独立自检")
    parser.add_argument("--version", action="version", version=f"{APP_NAME} {APP_VERSION}")
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.self_test:
            print(json.dumps(run_self_test(args.db if "--db" in (argv or sys.argv[1:]) else None), ensure_ascii=False, indent=2))
            return 0
        repo = EnergyRepository(args.db)
        if args.import_file:
            print(json.dumps(import_file(repo, args.import_file), ensure_ascii=False, indent=2))
        if args.export_file:
            count = export_file(repo, args.export_file)
            print(json.dumps({"ok": True, "exported": count, "path": str(args.export_file)}, ensure_ascii=False, indent=2))
        if args.no_gui or args.import_file or args.export_file:
            return 0
        import tkinter as tk
        root = tk.Tk()
        EnergyDailyApp(root, repo)
        root.mainloop()
        return 0
    except Exception as exc:
        print(f"运行失败：{friendly_message(exc)}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
