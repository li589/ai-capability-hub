"""Build the reusable DOCX shell for fission-distribution reports."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "report-template.docx"

FONT = "Microsoft YaHei"
BLUE = "0047AB"
DARK_BLUE = "00215E"
INK = "172033"
MUTED = "5D6778"
LIGHT_BLUE = "EAF2FF"
LIGHT_GRAY = "F2F4F7"
RED = "8B1E1E"
CONTENT_WIDTH_DXA = 9638  # A4 width minus 20 mm margins on both sides.


def _set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _set_style(style, size, color=INK, bold=False, before=0, after=6, line=1.1):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), FONT)
    rpr.rFonts.set(qn("w:hAnsi"), FONT)
    rpr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line


def _shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_fixed_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _cell_margins(cell)


def _page_field(paragraph):
    run = paragraph.add_run()
    _set_run_font(run, size=9, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def _configure_styles(doc):
    _set_style(doc.styles["Normal"], 10.5, after=6, line=1.1)
    _set_style(doc.styles["Title"], 28, color=BLUE, bold=True, after=8, line=1.0)
    _set_style(doc.styles["Heading 1"], 16, color=BLUE, bold=True, before=16, after=8)
    _set_style(doc.styles["Heading 2"], 13, color=DARK_BLUE, bold=True, before=12, after=6)
    _set_style(doc.styles["Heading 3"], 11.5, color=DARK_BLUE, bold=True, before=8, after=4)
    _set_style(doc.styles["Caption"], 9, color=MUTED, after=8, line=1.0)
    doc.styles["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    disclaimer = doc.styles.add_style("Distribution Disclaimer", WD_STYLE_TYPE.PARAGRAPH)
    _set_style(disclaimer, 9, color=RED, after=8, line=1.1)
    citation = doc.styles.add_style("Table Citation", WD_STYLE_TYPE.PARAGRAPH)
    _set_style(citation, 8.5, color=MUTED, before=4, after=4, line=1.0)


def _configure_section(section):
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.header_distance = Mm(12.5)
    section.footer_distance = Mm(12.5)
    section.different_first_page_header_footer = False


def _add_header_footer(section):
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(2)
    run = header.add_run("裂变式分销方案设计 · 内部决策版")
    _set_run_font(run, size=8.5, color=MUTED, bold=True)
    p_pr = header._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "D8E1EF")
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(2)
    run = footer.add_run("仅供内部商业策划参考  ·  第 ")
    _set_run_font(run, size=8.5, color=MUTED)
    _page_field(footer)
    run = footer.add_run(" 页")
    _set_run_font(run, size=8.5, color=MUTED)


def _add_cover(doc):
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    _set_run_font(kicker.add_run("合规优先 · 数据闭环 · 可执行落地"), size=10, color=BLUE, bold=True)
    title = doc.add_paragraph("裂变式分销方案", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(54)
    _set_run_font(subtitle.add_run("内部决策报告与外部招募物料模板"), size=14, color=DARK_BLUE)
    for label, value in (("项目名称", "[填写项目名称]"), ("企业 / 品牌", "[填写企业或品牌]"), ("报告日期", "[YYYY年MM月DD日]")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(7)
        _set_run_font(p.add_run(f"{label}："), size=10, color=MUTED, bold=True)
        _set_run_font(p.add_run(value), size=10.5, color=INK)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(54)
    _set_run_font(note.add_run("依据已确认事实、可追溯来源与封闭财务模型生成"), size=9, color=MUTED, italic=True)


def _add_body_sample(doc):
    doc.add_page_break()
    doc.add_heading("1. 执行摘要与项目目标", level=1)
    doc.add_paragraph("本页展示正式报告的正文层级、表格、图注与免责声明样式。生成实际项目报告时，应按十五节内部报告结构替换示例内容，并单独输出外部招募版本。")

    doc.add_heading("1.1 已确认事实与关键假设", level=2)
    table = doc.add_table(rows=4, cols=3)
    table.style = "Table Grid"
    _set_fixed_table_geometry(table, [1928, 4819, 2891])
    rows = [
        ("字段", "内容", "来源 / 定位"),
        ("产品 / 服务", "[填写已确认内容]", "[文件名·页/表/幻灯片]"),
        ("目标分销人群", "[填写已确认内容]", "[用户确认·日期]"),
        ("关键假设", "[写明区间、依据与日期]", "[外部权威来源]"),
    ]
    for row_i, values in enumerate(rows):
        for cell, value in zip(table.rows[row_i].cells, values):
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(value)
            _set_run_font(run, size=9.5, color=INK, bold=row_i == 0)
            if row_i == 0:
                _shade(cell, LIGHT_GRAY)

    doc.add_paragraph("表 1 来源：用户确认事实与注明日期的权威公开资料。", style="Table Citation")
    doc.add_heading("1.2 财务闭环示例", level=2)
    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    _set_fixed_table_geometry(callout, [CONTENT_WIDTH_DXA])
    _shade(callout.cell(0, 0), LIGHT_BLUE)
    p = callout.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    _set_run_font(p.add_run("收入 = 成本 + 直接成交佣金 + 运营费用 + 税费 + 剩余贡献毛利；所有金额必须闭合。"), size=10, color=DARK_BLUE, bold=True)

    doc.add_paragraph("[在此插入已渲染并验收的业务图表]", style="Caption")
    doc.add_paragraph("图 1：合规交易与直接佣金路径", style="Caption")
    doc.add_paragraph("本报告仅供商业策划参考，不构成法律、税务或监管意见；涉及受监管行业、宣传表述或奖励机制时，应由具备资质的专业人士复核。", style="Distribution Disclaimer")


def build():
    doc = Document()
    section = doc.sections[0]
    _configure_section(section)
    _configure_styles(doc)
    _add_header_footer(section)
    _add_cover(doc)
    _add_body_sample(doc)
    doc.core_properties.title = "裂变式分销方案报告模板"
    doc.core_properties.subject = "合规优先的分销方案内部报告模板"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
