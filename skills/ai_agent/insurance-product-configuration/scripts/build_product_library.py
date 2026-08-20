#!/usr/bin/env python3
"""Build a normalized insurance product catalog from the bundled product source files."""
from __future__ import annotations

import argparse
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

SOURCE_DATE = "2025-03-17"
VALIDITY = "有效（按代理人确认）"
MAIN_PRODUCTS = {"TANSPM001", "TLISPM001", "TLIM003", "THIM001", "TCIM001"}


def clean(value: Any) -> str:
    return "" if value is None else str(value).replace("\n", " ").strip()


def normalize_term(value: str) -> str:
    value = clean(value).replace("保险期限：", "").replace("保障至", "至")
    value = value.replace("保障终身", "终身")
    return value


def payment_label(value: str) -> str:
    value = clean(value).replace("交费", "交")
    if value in {"趸", "趸交", "一次性交清"}:
        return "趸交"
    if re.fullmatch(r"\d+年", value):
        return f"{value}交"
    return value


def parse_code_table(source: Path) -> dict[str, dict[str, str]]:
    workbook = load_workbook(source / "产品代码表.xlsx", data_only=True, read_only=True)
    sheet = workbook.active
    result: dict[str, dict[str, str]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        name, code, category = (clean(row[i]) if i < len(row) else "" for i in range(3))
        if code:
            result[code] = {"name": name, "category": category}
    result.setdefault("TCIA001", {"name": "附加守护星终身重疾险", "category": "重疾"})
    return result


def parse_combination_rules(source: Path, known_codes: set[str]) -> dict[str, list[dict[str, str]]]:
    workbook = load_workbook(source / "险种搭配规则.xlsx", data_only=True, read_only=True)
    sheet = workbook.active
    current_main = ""
    result: dict[str, list[dict[str, str]]] = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        main = clean(row[0]) if len(row) > 0 else ""
        rider = clean(row[1]) if len(row) > 1 else ""
        ratio = clean(row[2]) if len(row) > 2 else ""
        if main:
            current_main = main
        if current_main and rider:
            candidates = [code for code in known_codes if code.startswith(rider) or rider.startswith(code)]
            normalized_rider = candidates[0] if len(candidates) == 1 else rider
            result.setdefault(current_main, []).append({"rider_code": normalized_rider, "sum_assured_rule": ratio or "未标注"})
    return result


def parse_main_rules(source: Path) -> list[dict[str, str]]:
    document = Document(source / "主险产品基本投保规则.docx")
    if not document.tables:
        return []
    table = document.tables[0]
    rows: list[dict[str, str]] = []
    active_main = ""
    active_term = ""
    active_period = ""
    active_frequency = ""
    for cells in table.rows[1:]:
        values = [clean(cell.text) for cell in cells.cells]
        if values[0]:
            active_main = values[0]
        if values[1]:
            active_term = values[1]
        if values[3]:
            active_period = values[3]
        if values[4]:
            active_frequency = values[4]
        rows.append({
            "product_code": active_main,
            "payment_term": active_term,
            "issue_age": values[2],
            "coverage_term": active_period,
            "payment_frequency": active_frequency,
        })
    return rows


def source_files_for(code: str, source: Path) -> dict[str, str]:
    files = {"rate_table": "", "product_document": ""}
    for path in source.iterdir():
        if code in path.name:
            if "费率表" in path.name:
                files["rate_table"] = path.name
            elif "说明" in path.name:
                files["product_document"] = path.name
    if code == "TCIA001":
        files["rate_table"] = "附加守护星终身重疾险-产品费率表.xlsx"
        files["product_document"] = "附加守护星终身重疾险-产品说明书.pdf"
    return files


def extract_pdf_summary(path: Path | None) -> dict[str, str]:
    if path is None or not path.is_file():
        return {"document_id": "未标注", "waiting_period_excerpt": "待从有效条款核验", "exclusion_excerpt": "待从有效条款核验"}
    text = "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    identifier = re.search(r"HO-[A-Z]+-\d{4}-\d+", text)
    wait_at = text.find("等待期")
    exclusion_at = text.find("责任免除")
    return {
        "document_id": identifier.group(0) if identifier else "未标注",
        "waiting_period_excerpt": clean(text[wait_at:wait_at + 420]) if wait_at >= 0 else "待从有效条款核验",
        "exclusion_excerpt": clean(text[exclusion_at:exclusion_at + 520]) if exclusion_at >= 0 else "待从有效条款核验",
    }


def locate_gender_columns(sheet: Any, gender_row: int, payment_row: int) -> list[tuple[int, str, str]]:
    current_gender = ""
    columns: list[tuple[int, str, str]] = []
    for col in range(2, sheet.max_column + 1):
        gender_value = clean(sheet.cell(gender_row, col).value)
        if "男" in gender_value:
            current_gender = "男"
        elif "女" in gender_value:
            current_gender = "女"
        payment = payment_label(clean(sheet.cell(payment_row, col).value))
        if current_gender and payment:
            columns.append((col, current_gender, payment))
    return columns


def parse_per_thousand_basic(path: Path, code: str) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheets: dict[str, Any] = {}
    for sheet in workbook.worksheets:
        if "非标" in sheet.title:
            continue
        rows = list(sheet.iter_rows(min_row=1, max_row=6, values_only=True))
        gender_row = 3 if code in {"TCIM001", "TCIMUA002", "TCISCA003"} else 0
        payment_row = 4 if code in {"TCIM001", "TCIMUA002", "TCISCA003"} else 0
        if code == "TCIA001":
            term = "5年保证续保" if "5年" in sheet.title else "20年保证续保"
            rates: dict[str, Any] = {}
            for row in sheet.iter_rows(min_row=4, values_only=True):
                age = row[0] if row else None
                if isinstance(age, int):
                    rates[str(age)] = {"男": float(row[1]), "女": float(row[2])}
            sheets[term] = {"model": "per_thousand_sum_assured", "payment_term": "年交", "rates": rates}
            continue
        mappings = locate_gender_columns(sheet, gender_row, payment_row)
        rates: dict[str, dict[str, dict[str, float]]] = {}
        for row in sheet.iter_rows(min_row=payment_row + 1, values_only=True):
            age = row[0] if row else None
            if not isinstance(age, int):
                continue
            age_rates: dict[str, dict[str, float]] = {}
            for col, gender, payment in mappings:
                value = row[col - 1] if len(row) >= col else None
                if isinstance(value, (int, float)):
                    age_rates.setdefault(gender, {})[payment] = float(value)
            if age_rates:
                rates[str(age)] = age_rates
        term = normalize_term(sheet.title)
        sheets[term] = {"model": "per_thousand_sum_assured", "rates": rates}
    return {"product_code": code, "rate_type": "per_thousand_sum_assured", "sheets": sheets}


def parse_annuity(path: Path, code: str) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheets: dict[str, Any] = {}
    for sheet in workbook.worksheets:
        title = sheet.title
        gender = "男" if "男" in title else "女"
        collection_age = re.search(r"(\d+)周岁", title)
        headers = [payment_label(clean(sheet.cell(3, c).value)) for c in range(2, sheet.max_column + 1)]
        rates: dict[str, dict[str, float]] = {}
        for row in sheet.iter_rows(min_row=4, values_only=True):
            age = row[0] if row else None
            if not isinstance(age, int):
                continue
            values = {}
            for idx, term in enumerate(headers, start=1):
                value = row[idx] if len(row) > idx else None
                if term and isinstance(value, (int, float)):
                    values[term] = float(value)
            if values:
                rates[str(age)] = values
        key = f"{gender}-起始领取年龄{collection_age.group(1) if collection_age else '待确认'}"
        sheets[key] = {"model": "per_thousand_sum_assured", "rates": rates}
    return {"product_code": code, "rate_type": "per_thousand_sum_assured", "sheets": sheets}


def parse_life_premium_to_sum(path: Path, code: str) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    mappings = locate_gender_columns(sheet, 3, 4)
    rates: dict[str, dict[str, dict[str, float]]] = {}
    for row in sheet.iter_rows(min_row=5, values_only=True):
        age = row[0] if row else None
        if not isinstance(age, int):
            continue
        age_rates: dict[str, dict[str, float]] = {}
        for col, gender, payment in mappings:
            value = row[col - 1] if len(row) >= col else None
            if isinstance(value, (int, float)):
                age_rates.setdefault(gender, {})[payment] = float(value)
        if age_rates:
            rates[str(age)] = age_rates
    return {"product_code": code, "rate_type": "per_thousand_annual_premium_to_sum_assured", "sheets": {"标准": {"model": "per_thousand_annual_premium_to_sum_assured", "rates": rates}}}


def parse_medical(path: Path, code: str) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    bands = []
    for row in sheet.iter_rows(min_row=4, values_only=True):
        age_band = clean(row[0]) if row else ""
        social = row[1] if len(row) > 1 else None
        nonsocial = row[2] if len(row) > 2 else None
        if age_band and isinstance(social, (int, float)) and isinstance(nonsocial, (int, float)):
            bands.append({"age_band": age_band, "有社保": float(social), "无社保": float(nonsocial)})
    return {"product_code": code, "rate_type": "annual_fixed_by_age_band", "sheets": {"标准": {"model": "annual_fixed_by_age_band", "bands": bands}}}


def parse_term_life(path: Path, code: str) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheets: dict[str, Any] = {}
    for sheet in workbook.worksheets:
        gender = "男" if "男" in sheet.title else "女"
        coverage = ""
        mappings: list[tuple[int, str, str]] = []
        for col in range(2, sheet.max_column + 1):
            coverage_cell = clean(sheet.cell(2, col).value)
            if coverage_cell:
                coverage = normalize_term(coverage_cell)
            payment = payment_label(clean(sheet.cell(3, col).value))
            if coverage and payment:
                mappings.append((col, coverage, payment))
        rates: dict[str, dict[str, dict[str, float]]] = {}
        for row in sheet.iter_rows(min_row=4, values_only=True):
            age = row[0] if row else None
            if not isinstance(age, int):
                continue
            age_rates: dict[str, dict[str, float]] = {}
            for col, coverage_term, payment in mappings:
                value = row[col - 1] if len(row) >= col else None
                if isinstance(value, (int, float)):
                    age_rates.setdefault(coverage_term, {})[payment] = float(value)
            if age_rates:
                rates[str(age)] = age_rates
        sheets[gender] = {"model": "per_thousand_sum_assured", "rates": rates}
    return {"product_code": code, "rate_type": "per_thousand_sum_assured", "sheets": sheets}


def build_rates(source: Path, products: dict[str, dict[str, Any]]) -> dict[str, Any]:
    rates: dict[str, Any] = {}
    for code, product in products.items():
        rate_name = product["sources"].get("rate_table", "")
        if not rate_name:
            continue
        path = source / rate_name
        if code in {"TCIM001", "TCIMUA002", "TCISCA003", "TCIA001"}:
            rates[code] = parse_per_thousand_basic(path, code)
        elif code == "TANSPM001":
            rates[code] = parse_annuity(path, code)
        elif code in {"TLIM003", "TLISPM001"}:
            rates[code] = parse_life_premium_to_sum(path, code)
        elif code == "THIM001":
            rates[code] = parse_medical(path, code)
        elif code == "TLITLA001":
            rates[code] = parse_term_life(path, code)
        else:
            rates[code] = {"product_code": code, "rate_type": "needs_manual_rule", "sheets": {}}
    return rates


def main() -> None:
    parser = argparse.ArgumentParser(description="Build normalized product library JSON")
    parser.add_argument("--source", required=True, help="Directory containing product files")
    parser.add_argument("--output", required=True, help="Output catalog JSON path")
    parser.add_argument("--rates-output", required=True, help="Output rates JSON path")
    args = parser.parse_args()
    source = Path(args.source)
    code_map = parse_code_table(source)
    combination_rules = parse_combination_rules(source, set(code_map))
    main_rules = parse_main_rules(source)
    products: dict[str, dict[str, Any]] = {}
    for code, item in code_map.items():
        sources = source_files_for(code, source)
        pdf_path = source / sources["product_document"] if sources["product_document"] else None
        products[code] = {
            "code": code,
            "name": item["name"],
            "category": item["category"],
            "is_main": code in MAIN_PRODUCTS,
            "validity": VALIDITY,
            "library_reference_version": f"资料库版 {SOURCE_DATE}",
            "sources": sources,
            "terms": extract_pdf_summary(pdf_path),
        }
    catalog = {
        "library_name": "脱敏版产品库",
        "validity": VALIDITY,
        "library_reference_version": f"资料库版 {SOURCE_DATE}",
        "products": products,
        "combination_rules": combination_rules,
        "main_product_rules": main_rules,
        "source_manifest": [p.name for p in sorted(source.iterdir()) if p.is_file()],
        "built_at": datetime.now().isoformat(timespec="seconds"),
    }
    Path(args.output).write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    Path(args.rates_output).write_text(json.dumps(build_rates(source, products), ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
